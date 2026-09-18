"""Codex SDK runtime for durable Feishu assistant turns."""

from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutureTimeoutError
from datetime import date, datetime
import json
import logging
import math
from numbers import Real
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
from typing import Any, Callable, Optional, Tuple

from china_a_share.feishu_agent import (
    FeishuAgentOutcome,
    FeishuAgentRequest,
    FeishuResearchVisualization,
    MAX_AGENT_CONTEXT_TURNS,
)


logger = logging.getLogger(__name__)
CODEX_MODEL_PROVIDER = "deepseek"
CODEX_TURN_TIMEOUT_SECONDS = 900
CODEX_INTERRUPT_GRACE_SECONDS = 30
SUPPORTED_ARTIFACT_SUFFIXES = {".csv", ".docx", ".pdf", ".xlsx"}
MAX_VISUALIZATION_ROWS = 2_000
MAX_VISUALIZATION_COLUMNS = 24
MAX_ARTIFACT_FILENAME_STEM_LENGTH = 80
INVALID_ARTIFACT_FILENAME_PATTERN = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


class CodexFeishuAgentRuntime:
    """Run one Feishu turn with the full Codex agent harness."""

    def __init__(
        self,
        *,
        base_url: str,
        model: str,
        api_key: str,
        tushare_token: str,
        cache_bucket: str,
        sandbox_url: str,
        massive_api_key: str = "",
        finnhub_api_key: str = "",
        google_cloud_project: str = "",
        sdk_loader: Optional[Callable[[], Tuple[Any, Any, Any, Any]]] = None,
    ) -> None:
        self._base_url = base_url.rstrip("/") + "/"
        self._model = model
        self._api_key = api_key
        self._tushare_token = tushare_token
        self._massive_api_key = massive_api_key
        self._finnhub_api_key = finnhub_api_key
        self._cache_bucket = cache_bucket
        self._sandbox_url = sandbox_url
        self._google_cloud_project = google_cloud_project
        self._sdk_loader = sdk_loader or _load_codex_sdk

    def run(
        self,
        request: FeishuAgentRequest,
        progress: Callable[[str, str], None],
    ) -> FeishuAgentOutcome:
        """Execute the request without imposing an application-level turn limit."""
        Codex, CodexConfig, Sandbox, ApprovalMode = self._sdk_loader()
        with tempfile.TemporaryDirectory(prefix="feishu-codex-") as workspace_value:
            workspace = Path(workspace_value)
            artifact_dir = workspace / "artifacts"
            artifact_dir.mkdir()
            config = CodexConfig(
                cwd=str(workspace),
                env=self._codex_environment(artifact_dir, request),
                config_overrides=self._codex_overrides(),
            )
            progress(
                "researching",
                "Codex 正在调用通用工具并处理完整数据集…",
            )
            with Codex(config) as codex:
                thread = codex.thread_start(
                    approval_mode=ApprovalMode.deny_all,
                    cwd=str(workspace),
                    developer_instructions=_developer_instructions(),
                    ephemeral=True,
                    model=self._model,
                    model_provider=CODEX_MODEL_PROVIDER,
                    sandbox=Sandbox.workspace_write,
                )
                turn = thread.turn(_request_prompt(request))
                (
                    final_response,
                    duration_ms,
                    token_usage,
                    usage_update_count,
                ) = _run_turn_with_progress(turn, progress)

            artifact_path = _persist_artifact(
                artifact_dir,
                request.conversation_name,
            )
            visualization = _build_research_visualization(artifact_path)
            answer = str(final_response or "").strip()
            if not answer:
                answer = _empty_response_follow_up(artifact_path)
                logger.warning(
                    "codex_feishu_turn_empty_response conversation_id=%s model=%s "
                    "duration_ms=%s artifact=%s",
                    request.conversation_id,
                    self._model,
                    duration_ms,
                    artifact_path.name if artifact_path is not None else "none",
                )
            logger.info(
                "codex_feishu_turn_completed conversation_id=%s model=%s "
                "duration_ms=%s artifact=%s",
                request.conversation_id,
                self._model,
                duration_ms,
                artifact_path.name if artifact_path is not None else "none",
            )
            _log_turn_usage(
                request,
                self._model,
                token_usage,
                usage_update_count,
            )
            return FeishuAgentOutcome(
                answer=answer,
                artifact_path=artifact_path,
                visualization=visualization,
            )

    def _codex_environment(
        self,
        artifact_dir: Path,
        request: FeishuAgentRequest,
    ) -> dict[str, str]:
        env = {
            "LLM_API_KEY": self._api_key,
            "TUSHARE_TOKEN": self._tushare_token,
            "TUSHARE_CACHE_BUCKET": self._cache_bucket,
            "RESEARCH_SANDBOX_URL": self._sandbox_url,
            "CODEX_AGENT_ARTIFACT_DIR": str(artifact_dir),
            "CODEX_AGENT_CONVERSATION_ID": request.conversation_id,
        }
        task_id = os.getenv("ANALYSIS_TASK_ID", "").strip()
        if task_id:
            env["ANALYSIS_TASK_ID"] = task_id
        if self._google_cloud_project:
            env["GOOGLE_CLOUD_PROJECT"] = self._google_cloud_project
        if self._massive_api_key:
            env["MASSIVE_API_KEY"] = self._massive_api_key
        if self._finnhub_api_key:
            env["FINNHUB_API_KEY"] = self._finnhub_api_key
        return env

    def _codex_overrides(self) -> tuple[str, ...]:
        model_catalog = Path(__file__).with_name("deepseek_models.json")
        mcp_env_vars = [
            "TUSHARE_TOKEN",
            "TUSHARE_CACHE_BUCKET",
            "RESEARCH_SANDBOX_URL",
            "CODEX_AGENT_ARTIFACT_DIR",
            "CODEX_AGENT_CONVERSATION_ID",
        ]
        if os.getenv("ANALYSIS_TASK_ID", "").strip():
            mcp_env_vars.append("ANALYSIS_TASK_ID")
        if self._google_cloud_project:
            mcp_env_vars.append("GOOGLE_CLOUD_PROJECT")
        if self._massive_api_key:
            mcp_env_vars.append("MASSIVE_API_KEY")
        if self._finnhub_api_key:
            mcp_env_vars.append("FINNHUB_API_KEY")
        values = {
            "model": self._model,
            "model_provider": CODEX_MODEL_PROVIDER,
            "model_reasoning_effort": "high",
            "model_catalog_json": str(model_catalog),
            "forced_login_method": "api",
            "check_for_update_on_startup": False,
            "web_search": "disabled",
            "history.persistence": "none",
            "features.apps": False,
            "features.multi_agent": False,
            "features.remote_plugin": False,
            "features.hooks": False,
            "features.goals": False,
            "shell_environment_policy.inherit": "core",
            "shell_environment_policy.ignore_default_excludes": False,
            f"model_providers.{CODEX_MODEL_PROVIDER}.name": "DeepSeek",
            f"model_providers.{CODEX_MODEL_PROVIDER}.base_url": self._base_url,
            f"model_providers.{CODEX_MODEL_PROVIDER}.env_key": "LLM_API_KEY",
            f"model_providers.{CODEX_MODEL_PROVIDER}.wire_api": "responses",
            "mcp_servers.market_data.command": sys.executable,
            "mcp_servers.market_data.args": ["-m", "china_a_share.codex_mcp"],
            "mcp_servers.market_data.env_vars": mcp_env_vars,
            "mcp_servers.market_data.startup_timeout_sec": 30,
            "mcp_servers.market_data.tool_timeout_sec": 600,
            "mcp_servers.market_data.required": True,
            "mcp_servers.market_data.default_tools_approval_mode": "approve",
        }
        return tuple(
            f"{key}={json.dumps(value, ensure_ascii=True)}"
            for key, value in values.items()
        )


def _load_codex_sdk() -> Tuple[Any, Any, Any, Any]:
    try:
        from openai_codex import ApprovalMode, Codex, CodexConfig, Sandbox
    except ImportError as exc:
        raise RuntimeError(
            "The Feishu agent requires the openai-codex runtime dependency."
        ) from exc
    return Codex, CodexConfig, Sandbox, ApprovalMode


def _run_turn_with_progress(
    turn: Any,
    progress: Callable[[str, str], None],
) -> tuple[Optional[str], Optional[int], Optional[dict[str, int]], int]:
    """Collect one turn while surfacing material MCP activity and bounding runtime."""
    executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="codex-turn")
    future = executor.submit(_collect_turn, turn, progress)
    try:
        return future.result(timeout=CODEX_TURN_TIMEOUT_SECONDS)
    except FutureTimeoutError as exc:
        logger.error(
            "codex_feishu_turn_timeout timeout_seconds=%s",
            CODEX_TURN_TIMEOUT_SECONDS,
        )
        turn.interrupt()
        try:
            future.result(timeout=CODEX_INTERRUPT_GRACE_SECONDS)
        except Exception:
            logger.exception("codex_feishu_turn_interrupt_failed")
        raise RuntimeError(
            f"Codex turn exceeded {CODEX_TURN_TIMEOUT_SECONDS // 60} minutes."
        ) from exc
    finally:
        executor.shutdown(wait=False, cancel_futures=True)


def _collect_turn(
    turn: Any,
    progress: Callable[[str, str], None],
) -> tuple[Optional[str], Optional[int], Optional[dict[str, int]], int]:
    final_response: Optional[str] = None
    last_unknown_phase_response: Optional[str] = None
    duration_ms: Optional[int] = None
    token_usage: Optional[dict[str, int]] = None
    usage_update_count = 0
    completed = False
    for event in turn.stream():
        payload = _event_payload(event.payload)
        if event.method == "thread/tokenUsage/updated":
            # Token usage notifications are cumulative snapshots. Keep only the
            # latest total so repeated notifications are never double-counted.
            usage_update_count += 1
            token_usage = _normalize_token_usage(
                ((payload.get("tokenUsage") or {}).get("total") or {})
            ) or token_usage
            continue
        if event.method == "item/started":
            item = payload.get("item") or {}
            if item.get("type") == "mcpToolCall":
                tool = str(item.get("tool") or "unknown")
                progress("tool", _tool_progress_message(tool))
                logger.info("codex_mcp_tool_started tool=%s", tool)
            continue
        if event.method == "item/completed":
            item = payload.get("item") or {}
            item_type = item.get("type")
            if item_type == "mcpToolCall":
                logger.info(
                    "codex_mcp_tool_completed tool=%s status=%s duration_ms=%s",
                    item.get("tool"),
                    item.get("status"),
                    item.get("durationMs"),
                )
            elif item_type == "agentMessage":
                text = str(item.get("text") or "").strip()
                if not text:
                    continue
                if item.get("phase") == "final_answer":
                    final_response = text
                elif item.get("phase") is None:
                    # The official SDK accepts the latest agent message whose phase
                    # is absent when a model provider does not emit final-answer
                    # metadata. DeepSeek uses this valid compatibility path.
                    last_unknown_phase_response = text
            continue
        if event.method == "turn/completed":
            completed = True
            completed_turn = payload.get("turn") or {}
            duration_ms = completed_turn.get("durationMs")
            status = str(completed_turn.get("status") or "")
            if status != "completed":
                error = completed_turn.get("error") or {}
                message = error.get("message") or f"Codex turn ended with {status}."
                raise RuntimeError(str(message))
    if not completed:
        raise RuntimeError("Codex turn ended without a completion event.")
    return (
        final_response or last_unknown_phase_response,
        duration_ms,
        token_usage,
        usage_update_count,
    )


def _normalize_token_usage(value: Any) -> Optional[dict[str, int]]:
    """Return the stable raw Codex token counters from one usage snapshot."""
    if not isinstance(value, dict):
        return None
    fields = {
        "input_tokens": "inputTokens",
        "cached_input_tokens": "cachedInputTokens",
        "output_tokens": "outputTokens",
        "reasoning_output_tokens": "reasoningOutputTokens",
        "total_tokens": "totalTokens",
    }
    normalized: dict[str, int] = {}
    for output_name, source_name in fields.items():
        raw_value = value.get(source_name)
        if isinstance(raw_value, bool) or not isinstance(raw_value, Real):
            return None
        normalized[output_name] = max(int(raw_value), 0)
    return normalized


def _log_turn_usage(
    request: FeishuAgentRequest,
    model: str,
    token_usage: Optional[dict[str, int]],
    usage_update_count: int,
) -> None:
    """Log task-scoped raw usage without coupling runtime behavior to billing rates."""
    usage = token_usage or {}
    workload = (
        "live_regression"
        if os.getenv("RUN_LIVE_ANALYSIS", "").strip() == "1"
        else "feishu_user"
    )
    logger.info(
        "codex_feishu_turn_usage task_id=%s conversation_id=%s model=%s "
        "workload=%s usage_available=%s usage_updates=%s input_tokens=%s "
        "cached_input_tokens=%s output_tokens=%s reasoning_output_tokens=%s "
        "total_tokens=%s",
        os.getenv("ANALYSIS_TASK_ID", "").strip() or "none",
        request.conversation_id,
        model,
        workload,
        bool(token_usage),
        usage_update_count,
        usage.get("input_tokens", 0),
        usage.get("cached_input_tokens", 0),
        usage.get("output_tokens", 0),
        usage.get("reasoning_output_tokens", 0),
        usage.get("total_tokens", 0),
    )


def _empty_response_follow_up(artifact_path: Optional[Path]) -> str:
    """Return a recoverable user response when a completed turn has no text."""
    if artifact_path is not None:
        return "研究结果文件已生成，请查看附件。"
    return (
        "这次没有生成可验证的回答。请选择下一步：\n"
        "1. 按原问题重试（推荐）\n"
        "2. 补充查询范围、时间和指标口径后再试\n"
        "请回复序号，或直接补充你的完整口径。"
    )


def _event_payload(payload: Any) -> dict[str, Any]:
    if hasattr(payload, "model_dump"):
        value = payload.model_dump(
            by_alias=True,
            exclude_none=True,
            mode="json",
        )
        if isinstance(value, dict):
            return value
    if isinstance(payload, dict):
        return payload
    return {}


def _tool_progress_message(tool: str) -> str:
    messages = {
        "inspect_session_dataset": "Codex 正在读取当前会话的完整结果…",
        "request_clarification": "Codex 正在整理需要你确认的选项…",
        "search_market_data": "Codex 正在查找可用数据接口…",
        "query_market_data": "Codex 正在读取完整数据集…",
        "rank_dataset": "Codex 正在排序完整数据集…",
        "join_datasets": "Codex 正在关联完整数据集…",
        "transform_dataset": "Codex 正在转换完整数据集…",
        "run_python_analysis": "Codex 正在沙箱中计算完整数据集…",
        "export_excel": "Codex 正在生成 Excel 文件…",
    }
    return messages.get(tool, "Codex 正在执行通用工具…")


def _developer_instructions() -> str:
    return (
        "You are the remote assistant behind a Feishu conversation. Answer the "
        "current user in concise Chinese. Use the available MCP data tools whenever "
        "the answer depends on external or current structured data. Treat every tool "
        "preview as display-only: retained-dataset tools and the Python sandbox work "
        "with the complete dataset identified by dataset_id. Prefer generic query, "
        "join, transform, rank, and sandbox composition over assumptions or manual "
        "reconstruction. If a material product or analytical choice is ambiguous, "
        "call request_clarification once with two to four numbered choices, mark the "
        "safest default as recommended, and return its clarification verbatim. When "
        "conversation history shows the user answering that clarification, resolve "
        "the answer from context instead of asking again. If the "
        "inspect_session_dataset tool is available and the user refers to the "
        "previous list, result, table, or screening output, inspect and reuse that "
        "complete session dataset instead of reconstructing it from text or querying "
        "the same base universe again. For calculations that need "
        "several tabular operations, "
        "prefer one Python sandbox call over a long sequence of transformations. If "
        "several independent data reads are required, request them in the same model "
        "turn. Search the operation catalog only when the operation or parameter "
        "shape is not already established, and never repeat an equivalent catalog "
        "search or provider query within one turn. "
        "a tool rejects invalid arguments, inspect its schema and correct the call "
        "once; do not repeat equivalent failing calls. Never invent missing values. "
        "If a result has more than ten rows or the user asks for a file, create an "
        "Excel artifact with the export tool. When the user asks for a chart or "
        "interactive visualization, export the complete final dataset to Excel so "
        "the delivery layer can render it interactively. Do not create keyword routes, "
        "domain-specific shortcuts, or special cases for individual questions."
    )


def _request_prompt(request: FeishuAgentRequest) -> str:
    conversation = [
        {"user": turn.prompt, "assistant": turn.answer}
        for turn in request.conversation[-MAX_AGENT_CONTEXT_TURNS:]
    ]
    payload = {
        "previous_conversation": conversation,
        "current_user_request": request.prompt,
    }
    return (
        "Continue the following backend-managed conversation. Previous exchanges are "
        "context only; answer current_user_request directly.\n"
        + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    )


def _persist_artifact(
    artifact_dir: Path,
    conversation_name: str,
) -> Optional[Path]:
    candidates = [
        path
        for path in artifact_dir.iterdir()
        if path.is_file()
        and not path.is_symlink()
        and path.suffix.casefold() in SUPPORTED_ARTIFACT_SUFFIXES
        and path.stat().st_size > 0
    ]
    if not candidates:
        return None
    if len(candidates) > 1:
        raise RuntimeError("Codex produced more than one terminal artifact.")
    source = candidates[0]
    output_dir = Path(tempfile.mkdtemp(prefix="feishu-agent-output-"))
    output_path = output_dir / (
        _safe_artifact_filename_stem(conversation_name) + source.suffix.casefold()
    )
    shutil.copy2(source, output_path)
    return output_path


def _safe_artifact_filename_stem(value: str) -> str:
    """Return a portable filename stem while preserving the session name."""
    sanitized = INVALID_ARTIFACT_FILENAME_PATTERN.sub("_", value).strip(" .")
    sanitized = sanitized[:MAX_ARTIFACT_FILENAME_STEM_LENGTH].rstrip(" .")
    return sanitized or "a_share_research"


def _build_research_visualization(
    artifact_path: Optional[Path],
) -> Optional[FeishuResearchVisualization]:
    """Extract a bounded JSON-safe viewer dataset from one generated workbook."""
    if artifact_path is None or artifact_path.suffix.casefold() != ".xlsx":
        return None
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(artifact_path, read_only=True, data_only=True)
        sheet = workbook["Results"] if "Results" in workbook.sheetnames else workbook.active
        title = str(sheet["A2"].value or "A股研究结果").strip()
        raw_headers = next(
            sheet.iter_rows(min_row=5, max_row=5, values_only=True),
            (),
        )
        headers = [
            str(value).strip()
            for value in raw_headers[:MAX_VISUALIZATION_COLUMNS]
            if value is not None and str(value).strip()
        ]
        if not headers:
            workbook.close()
            return None
        rows = []
        for values in sheet.iter_rows(
            min_row=6,
            max_col=len(headers),
            values_only=True,
        ):
            if len(rows) >= MAX_VISUALIZATION_ROWS:
                break
            if not any(value is not None for value in values):
                continue
            rows.append(
                {
                    column: _json_safe_cell(value)
                    for column, value in zip(headers, values)
                }
            )
        source_row_count = max(sheet.max_row - 5, 0)
        workbook.close()
        numeric_columns = [
            column
            for column in headers
            if any(_is_numeric(row.get(column)) for row in rows)
        ]
        suggested_x = next(
            (column for column in headers if column not in numeric_columns),
            headers[0],
        )
        suggested_y = numeric_columns[-1] if numeric_columns else ""
        return FeishuResearchVisualization(
            title=title,
            columns=headers,
            numeric_columns=numeric_columns,
            rows=rows,
            source_row_count=source_row_count,
            truncated=source_row_count > len(rows),
            suggested_x=suggested_x,
            suggested_y=suggested_y,
        )
    except Exception:
        # Workbook delivery remains useful even when a malformed sheet cannot be
        # converted into the optional interactive viewer contract.
        logger.exception(
            "codex_feishu_visualization_extract_failed artifact=%s",
            artifact_path.name,
        )
        return None


def _json_safe_cell(value: Any) -> Any:
    """Normalize workbook scalars into bounded JSON-compatible values."""
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Real):
        number = float(value)
        return number if math.isfinite(number) else None
    return str(value)


def _is_numeric(value: Any) -> bool:
    """Return whether one normalized viewer value is a finite number."""
    return isinstance(value, (int, float)) and not isinstance(value, bool)
