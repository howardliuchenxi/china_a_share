"""Codex-style tool-using research runtime for the Feishu channel."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
import hashlib
import logging
from pathlib import Path
import re
import secrets
import tempfile
from typing import Any, Callable, Dict, List, Literal, Mapping, Optional, Protocol
from urllib.parse import urlencode
from uuid import uuid4
from zoneinfo import ZoneInfo

import pandas as pd
from pydantic import BaseModel, ConfigDict, Field

from china_a_share.core.contracts import (
    MAX_ANALYSIS_PROMPT_LENGTH,
    AnalysisTaskStatus,
    QueryResult,
    QueryStatus,
    ResultPipeline,
    ServiceError,
)
from china_a_share.result_pipeline import ResultPipelineExecutor
from china_a_share.registry import data_recency_note, field_unit_notes_for
from china_a_share.research_manifest import (
    ResearchManifest,
    finalize_research_manifest,
    new_research_manifest,
    record_dataset_evidence,
    research_dataset_evidence,
    select_final_dataset,
)
from china_a_share.security_links import (
    is_security_code_column,
    security_quote_page_url,
)


MAX_AGENT_PREVIEW_ROWS = 10
MAX_AGENT_CONTEXT_TURNS = 12
MAX_AGENT_PROGRESS_UPDATES = 19
RESEARCH_VISUALIZATION_LINK_LIFETIME = timedelta(days=30)
MIN_COLUMN_NOTE_CHARACTERS = 8
MAX_COLUMN_NOTE_CHARACTERS = 400
MAX_VISUALIZATION_METHODOLOGY_CHARACTERS = 6_000
AGENT_EXECUTION_LEASE_DURATION = timedelta(hours=8)
MAX_TERMINAL_DELIVERY_ATTEMPTS = 2
logger = logging.getLogger(__name__)


# Feishu application code 230011: the replied-to source message was withdrawn
# by its author, so every further reply or file upload to it will keep failing.
# Defined here because feishu.py imports its shared models from this module;
# the Feishu client raises it and the coordinator consumes it.
FEISHU_MESSAGE_WITHDRAWN_CODE = 230011


class FeishuSourceMessageWithdrawnError(RuntimeError):
    """Report that the source message being replied to no longer exists."""


class _FeishuSourceWithdrawnCancellation(Exception):
    """Abort one task whose triggering Feishu message was already withdrawn."""


class FeishuAgentConversationTurn(BaseModel):
    """One completed user and assistant exchange retained by a Feishu session."""

    model_config = ConfigDict(extra="forbid")

    prompt: str = Field(min_length=1, max_length=MAX_ANALYSIS_PROMPT_LENGTH)
    answer: str = Field(min_length=1, max_length=12_000)


class FeishuAgentRequest(BaseModel):
    """Immutable input required to execute one Feishu agent turn."""

    model_config = ConfigDict(extra="forbid")

    prompt: str = Field(min_length=1, max_length=MAX_ANALYSIS_PROMPT_LENGTH)
    conversation_id: str = Field(min_length=1)
    conversation_name: str = Field(
        default="默认会话",
        min_length=1,
        max_length=80,
        description="User-visible session name used for delivered artifact filenames.",
    )
    source_message_id: str = Field(min_length=1)
    conversation: List[FeishuAgentConversationTurn] = Field(
        default_factory=list,
        max_length=MAX_AGENT_CONTEXT_TURNS,
    )


class FeishuResearchVisualization(BaseModel):
    """Bounded tabular data used by the public read-only research viewer."""

    model_config = ConfigDict(extra="forbid")

    title: str = Field(min_length=1, description="Human-readable research title.")
    columns: List[str] = Field(description="Ordered columns available to the viewer.")
    numeric_columns: List[str] = Field(
        description="Columns whose retained values are numeric."
    )
    rows: List[Dict[str, Any]] = Field(
        description="Bounded JSON-safe rows retained for interactive exploration."
    )
    source_row_count: int = Field(
        ge=0,
        description="Complete row count in the source workbook before viewer limits.",
    )
    truncated: bool = Field(
        description="Whether the viewer contains fewer rows than the source workbook."
    )
    suggested_x: str = Field(description="Default horizontal-axis column.")
    suggested_y: str = Field(description="Default numeric vertical-axis column.")
    column_notes: Dict[str, str] = Field(
        default_factory=dict,
        description=(
            "One concise reader-facing explanation per result column, extracted "
            "from the workbook column-notes sheet."
        ),
    )
    methodology: str = Field(
        default="",
        max_length=MAX_VISUALIZATION_METHODOLOGY_CHARACTERS,
        description=(
            "Bounded methodology text recorded with the workbook and shown by "
            "the read-only viewer."
        ),
    )


class FeishuDeliveryRecord(BaseModel):
    """One durable terminal delivery intent and its latest outcome."""

    model_config = ConfigDict(extra="forbid")

    delivery_id: str = Field(
        min_length=1,
        description="Stable task-local identifier used to de-duplicate delivery.",
    )
    kind: Literal["text_reply", "text_update", "artifact"] = Field(
        description="Feishu operation required to deliver the terminal content.",
    )
    status: Literal["pending", "sent", "failed", "skipped"] = Field(
        description="Latest durable state of this delivery intent.",
    )
    target_message_id: str = Field(
        min_length=1,
        description="Source or application-authored message targeted by the delivery.",
    )
    content_sha256: str = Field(
        min_length=64,
        max_length=64,
        description="SHA-256 digest that prevents an identifier from changing content.",
    )
    text: Optional[str] = Field(
        default=None,
        description="Non-secret semantic text retained for retry without recomputation.",
    )
    render_mode: Literal["plain", "terminal_with_visualization"] = Field(
        default="plain",
        description="Deterministic rendering mode applied when a text retry is sent.",
    )
    artifact_name: Optional[str] = Field(
        default=None,
        description="Persisted task artifact retried by this delivery intent.",
    )
    remote_message_id: Optional[str] = Field(
        default=None,
        description="Feishu message identifier returned for a successful reply.",
    )
    attempt_count: int = Field(
        default=0,
        ge=0,
        description="Number of external delivery attempts already made.",
    )
    last_error: Optional[str] = Field(
        default=None,
        description="Most recent bounded delivery error for operator diagnosis.",
    )
    created_at: datetime = Field(description="UTC time when the intent was created.")
    updated_at: datetime = Field(description="UTC time of the latest delivery change.")
    delivered_at: Optional[datetime] = Field(
        default=None,
        description="UTC time when Feishu acknowledged successful delivery.",
    )


class FeishuAgentTask(BaseModel):
    """Durable lifecycle and output for one independent Feishu agent turn."""

    model_config = ConfigDict(extra="forbid")

    task_type: Literal["feishu_agent"] = "feishu_agent"
    task_id: str = Field(min_length=1)
    status: AnalysisTaskStatus
    request: FeishuAgentRequest
    created_at: datetime
    updated_at: datetime
    stage: str = Field(default="queued", min_length=1)
    progress_message: str = Field(default="等待研究任务启动。", min_length=1)
    answer: Optional[str] = None
    artifact_name: Optional[str] = None
    visualization: Optional[FeishuResearchVisualization] = Field(
        default=None,
        description="Bounded dataset exposed by the token-protected research viewer.",
    )
    visualization_token_hash: Optional[str] = Field(
        default=None,
        description="SHA-256 digest of the bearer token required by the viewer.",
    )
    visualization_expires_at: Optional[datetime] = Field(
        default=None,
        description="UTC time after which the viewer and workbook are unavailable.",
    )
    execution_lease_owner: Optional[str] = Field(
        default=None,
        description="Opaque worker identifier holding the current execution lease.",
    )
    execution_lease_expires_at: Optional[datetime] = Field(
        default=None,
        description="UTC expiry after which an interrupted execution may be reclaimed.",
    )
    execution_attempt_count: int = Field(
        default=0,
        ge=0,
        description="Number of workers that successfully claimed task execution.",
    )
    progress_message_id: Optional[str] = Field(
        default=None,
        description="Application-authored Feishu message reused for progress updates.",
    )
    progress_update_count: int = Field(
        default=0,
        ge=0,
        description="Number of persisted updates made to the progress message.",
    )
    research_manifest: Optional[ResearchManifest] = Field(
        default=None,
        description="Machine-readable evidence and output fingerprints for this task.",
    )
    result_fingerprint: Optional[str] = Field(
        default=None,
        description="Stable digest of the request, evidence graph, and terminal outputs.",
    )
    delivery_outbox: List[FeishuDeliveryRecord] = Field(
        default_factory=list,
        description="Durable terminal text and artifact delivery intents.",
    )
    error: Optional[ServiceError] = None


class AgentTaskStore(Protocol):
    """Persist complete Feishu agent task records."""

    def get(self, task_id: str) -> Optional[Any]:
        """Return one task when it exists."""

    def put(self, task: FeishuAgentTask) -> None:
        """Create or replace one task."""

    def claim_feishu_task(
        self,
        task_id: str,
        lease_owner: str,
        lease_expires_at: datetime,
    ) -> Optional[FeishuAgentTask]:
        """Atomically claim queued or expired Feishu task execution."""

    def put_claimed_feishu_task(
        self,
        task: FeishuAgentTask,
        lease_owner: str,
        lease_expires_at: datetime,
    ) -> None:
        """Persist one task only while the caller still owns its lease."""

    def release_feishu_task(self, task: FeishuAgentTask, lease_owner: str) -> None:
        """Persist terminal state and clear the caller's execution lease."""

    def put_artifact(self, task_id: str, path: Path) -> None:
        """Persist one generated artifact under its task identifier."""

    def get_artifact(self, task_id: str, artifact_name: str) -> Optional[bytes]:
        """Return one persisted artifact when it exists."""

    def archive_dataset(self, task_id: str, result: QueryResult) -> None:
        """Persist one complete intermediate or final dataset for a task."""

    def mark_final_dataset(self, task_id: str, dataset_id: str) -> None:
        """Mark the dataset selected for the task's terminal artifact."""

    def promote_session_workspace(self, conversation_id: str, task_id: str) -> bool:
        """Promote the task's final or latest dataset into the session workspace."""

    def get_session_workspace(self, conversation_id: str) -> Optional[QueryResult]:
        """Return the complete dataset currently active for one named session."""

    def get_research_manifest(self, task_id: str) -> Optional[ResearchManifest]:
        """Return incremental research evidence retained for one task."""

    def put_research_manifest(self, manifest: ResearchManifest) -> None:
        """Persist incremental research evidence independently of task progress."""


class AgentTaskDispatcher(Protocol):
    """Dispatch one persisted Feishu agent task to a worker."""

    def dispatch(self, task_id: str) -> None:
        """Start asynchronous execution for one task."""


class AgentProgressSink(Protocol):
    """Deliver stage changes and terminal artifacts to Feishu."""

    def reply(self, message_id: str, text: str) -> str:
        """Reply with one progress message and return its message ID."""

    def update(self, message_id: str, text: str) -> None:
        """Replace a previously sent progress message."""

    def reply_file(self, message_id: str, path: Path) -> None:
        """Upload and reply with one generated file."""


class FeishuAgentRunner(Protocol):
    """Execute one backend-managed Feishu conversation turn."""

    def run(
        self,
        request: FeishuAgentRequest,
        progress: Callable[[str, str], None],
    ) -> "FeishuAgentOutcome":
        """Return one terminal answer and optional artifact."""


class MarketDataProvider(Protocol):
    """Expose the read-only provider surface available to the agent tools."""

    @property
    def name(self) -> str:
        """Return the provider identifier."""

    def search_operations(self, prompt: str) -> Any:
        """Return matching read-only provider operations."""

    def supports(self, operation: str) -> bool:
        """Return whether one operation is allowlisted."""

    def describe_query_shapes(self, operation: str) -> List[Dict[str, Any]]:
        """Return audited parameter shapes exposed to the research model."""

    def validate_query(
        self,
        operation: str,
        params: Dict[str, Any],
        fields: List[str],
    ) -> None:
        """Validate one provider-native request before network access."""

    def query(
        self,
        operation: str,
        params: Dict[str, Any],
        fields: List[str],
        *,
        api_route: str,
        request_id: str,
        query_id: str,
    ) -> pd.DataFrame:
        """Execute one audited read-only provider query."""


class PythonSandbox(Protocol):
    """Execute restricted DataFrame programs outside the credentialed worker."""

    def run(self, code: str, datasets: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
        """Return one validated tabular result from isolated execution."""


class FeishuAgentCoordinator:
    """Submit independent Feishu turns without conversation-level serialization."""

    def __init__(
        self,
        store: AgentTaskStore,
        dispatcher: AgentTaskDispatcher,
        *,
        public_app_url: str = "",
    ) -> None:
        self._store = store
        self._dispatcher = dispatcher
        self._public_app_url = public_app_url.rstrip("/")

    def submit(
        self,
        request: FeishuAgentRequest,
        *,
        task_id: Optional[str] = None,
    ) -> FeishuAgentTask:
        """Persist and dispatch an idempotent agent task."""
        task_id = task_id or uuid4().hex
        existing = self._store.get(task_id)
        if existing is not None:
            if not isinstance(existing, FeishuAgentTask) or existing.request != request:
                raise ValueError(f"Agent task identifier is already in use: {task_id}")
            return existing
        now = datetime.now(timezone.utc)
        task = FeishuAgentTask(
            task_id=task_id,
            status=AnalysisTaskStatus.QUEUED,
            request=request,
            created_at=now,
            updated_at=now,
        )
        self._store.put(task)
        try:
            self._dispatcher.dispatch(task_id)
        except Exception as exc:
            logger.exception("feishu_agent_dispatch_failed task_id=%s", task_id)
            task.status = AnalysisTaskStatus.FAILED
            task.stage = "failed"
            task.progress_message = "研究任务未能启动。"
            task.error = ServiceError(source="system", message=str(exc))
            task.updated_at = datetime.now(timezone.utc)
            self._store.put(task)
            raise
        return task

    def get(self, task_id: str) -> Optional[FeishuAgentTask]:
        """Return one Feishu agent task when it exists."""
        task = self._store.get(task_id)
        return task if isinstance(task, FeishuAgentTask) else None

    def run(
        self,
        task_id: str,
        runtime: FeishuAgentRunner,
        progress_sink: AgentProgressSink,
    ) -> FeishuAgentTask:
        """Execute one atomically claimed task and persist terminal deliveries."""
        task = self.get(task_id)
        if task is None:
            raise KeyError(f"Feishu agent task does not exist: {task_id}")
        lease_owner = uuid4().hex
        claimed = self._store.claim_feishu_task(
            task_id,
            lease_owner,
            datetime.now(timezone.utc) + AGENT_EXECUTION_LEASE_DURATION,
        )
        if claimed is None:
            # Another worker owns a live lease, or the task has already reached a
            # terminal state. Returning the latest record makes duplicate Cloud
            # Run dispatches harmless without changing the worker contract.
            latest = self.get(task_id)
            if latest is None:
                raise KeyError(f"Feishu agent task does not exist: {task_id}")
            return latest
        task = claimed

        def persist_claimed() -> None:
            self._store.put_claimed_feishu_task(
                task,
                lease_owner,
                datetime.now(timezone.utc) + AGENT_EXECUTION_LEASE_DURATION,
            )

        if task.status == AnalysisTaskStatus.SUCCEEDED:
            try:
                return self._resume_terminal_deliveries(
                    task,
                    progress_sink,
                    persist_claimed,
                )
            finally:
                self._store.release_feishu_task(task, lease_owner)

        last_notified_progress: Optional[tuple[str, str]] = None
        progress_message_id = task.progress_message_id
        source_withdrawn = False

        def publish_progress(message: str) -> None:
            nonlocal progress_message_id, source_withdrawn
            if source_withdrawn:
                return
            try:
                if progress_message_id is None:
                    progress_message_id = progress_sink.reply(
                        task.request.source_message_id,
                        message,
                    )
                    task.progress_message_id = progress_message_id
                    persist_claimed()
                    return
                # Feishu limits edits per message. Reserve the final permitted edit
                # for the durable terminal outbox while retaining task progress.
                if task.progress_update_count >= MAX_AGENT_PROGRESS_UPDATES:
                    return
                progress_sink.update(progress_message_id, message)
                task.progress_update_count += 1
                persist_claimed()
            except FeishuSourceMessageWithdrawnError:
                # Once the source message is withdrawn every later delivery to it
                # fails too; stop attempting and let the run decide to continue
                # or cancel based on how far it already progressed.
                source_withdrawn = True
                logger.info(
                    "feishu_agent_source_message_withdrawn task_id=%s", task_id
                )

        def report(stage: str, message: str) -> None:
            nonlocal last_notified_progress
            task.stage = stage
            task.progress_message = message
            task.updated_at = datetime.now(timezone.utc)
            persist_claimed()
            progress = (stage, message)
            # Repeated tool calls may emit the same status, so suppress only exact
            # consecutive duplicates while preserving distinct progress details.
            if progress == last_notified_progress:
                return
            publish_progress(message)
            last_notified_progress = progress

        try:
            report("planning", "正在理解问题并选择研究工具…")
            if source_withdrawn:
                # The user withdrew the triggering message before any research
                # started, so further work would bill the model invisibly.
                raise _FeishuSourceWithdrawnCancellation()
            outcome = runtime.run(task.request, report)
            task.answer = outcome.answer
            task.artifact_name = (
                outcome.artifact_path.name if outcome.artifact_path is not None else None
            )
            terminal_message = outcome.answer
            artifact_persisted = False
            if outcome.artifact_path is not None:
                try:
                    self._store.put_artifact(task_id, outcome.artifact_path)
                    artifact_persisted = True
                except Exception:
                    # Attachment delivery may still succeed when durable artifact
                    # storage is temporarily unavailable, so preserve the answer.
                    logger.exception(
                        "feishu_agent_artifact_archive_failed task_id=%s artifact=%s",
                        task_id,
                        outcome.artifact_path.name,
                    )
            if (
                outcome.visualization is not None
                and outcome.artifact_path is not None
                and artifact_persisted
                and self._public_app_url
            ):
                try:
                    token = secrets.token_urlsafe(32)
                    task.visualization = outcome.visualization
                    task.visualization_token_hash = research_visualization_token_hash(
                        token
                    )
                    task.visualization_expires_at = (
                        datetime.now(timezone.utc)
                        + RESEARCH_VISUALIZATION_LINK_LIFETIME
                    )
                    query = urlencode({"token": token})
                    visualization_url = (
                        f"{self._public_app_url}/research/{task_id}?{query}"
                    )
                    terminal_message += (
                        "\n\n研究结果页面（30天内有效，点击后直接查看）：\n"
                        + visualization_url
                    )
                except Exception:
                    # A viewer publication failure must not invalidate research or
                    # the separately delivered Feishu workbook attachment.
                    logger.exception(
                        "feishu_agent_visualization_publish_failed task_id=%s",
                        task_id,
                    )
                    task.visualization = None
                    task.visualization_token_hash = None
                    task.visualization_expires_at = None
            if self._store.promote_session_workspace(
                task.request.conversation_id,
                task_id,
            ):
                logger.info(
                    "feishu_agent_session_workspace_promoted task_id=%s "
                    "conversation_id=%s",
                    task_id,
                    task.request.conversation_id,
                )

            manifest = self._store.get_research_manifest(task_id)
            if manifest is None:
                manifest = new_research_manifest(task_id)
            manifest = finalize_research_manifest(
                manifest,
                request=task.request.prompt,
                answer=outcome.answer,
                artifact_path=outcome.artifact_path,
            )
            self._store.put_research_manifest(manifest)
            task.research_manifest = manifest
            task.result_fingerprint = manifest.output_fingerprint
            task.status = AnalysisTaskStatus.SUCCEEDED
            task.stage = "completed"
            task.progress_message = "研究完成。"
            task.updated_at = datetime.now(timezone.utc)
            # Queue every terminal side effect in the same durable task write as
            # success. A worker crash after this point can therefore retry delivery
            # without either losing the intent or rerunning completed research.
            self._queue_terminal_text(task, terminal_message)
            if outcome.artifact_path is not None and outcome.artifact_path.exists():
                self._queue_artifact_record(
                    task,
                    outcome.artifact_path.read_bytes(),
                )
            persist_claimed()

            terminal_delivered = self._deliver_terminal_text(
                task,
                terminal_message,
                progress_sink,
                persist_claimed,
                source_withdrawn=source_withdrawn,
            )
            artifact_delivered = True
            if outcome.artifact_path is not None:
                artifact_delivered = self._deliver_terminal_artifact(
                    task,
                    progress_sink,
                    persist_claimed,
                    artifact_path=outcome.artifact_path,
                    source_withdrawn=source_withdrawn,
                )
            if terminal_delivered and not artifact_delivered and not source_withdrawn:
                task.progress_message = "研究完成，但附件发送失败。"
                task.updated_at = datetime.now(timezone.utc)
                persist_claimed()
                self._deliver_text_record(
                    task,
                    delivery_id="artifact-failure-notice",
                    text=(
                        (task.answer or outcome.answer)
                        + "\n\n研究已完成，但附件发送失败。请稍后回复“重试”重新生成附件。"
                    ),
                    progress_sink=progress_sink,
                    persist=persist_claimed,
                    source_withdrawn=False,
                )
        except _FeishuSourceWithdrawnCancellation:
            logger.info(
                "feishu_agent_cancelled_source_withdrawn task_id=%s", task_id
            )
            task.status = AnalysisTaskStatus.FAILED
            task.stage = "cancelled"
            task.progress_message = "提问消息已撤回，研究已取消。"
            task.error = ServiceError(
                source="system",
                message="Feishu source message was withdrawn; task cancelled.",
            )
            task.updated_at = datetime.now(timezone.utc)
            persist_claimed()
        except Exception as exc:
            logger.exception("feishu_agent_execution_failed task_id=%s", task_id)
            task.status = AnalysisTaskStatus.FAILED
            task.stage = "failed"
            task.progress_message = "研究任务失败。"
            task.error = ServiceError(source="system", message=str(exc))
            task.updated_at = datetime.now(timezone.utc)
            if progress_message_id is not None and not source_withdrawn:
                self._deliver_text_record(
                    task,
                    delivery_id="terminal-failure",
                    text=f"研究任务失败：{exc}",
                    progress_sink=progress_sink,
                    persist=persist_claimed,
                    source_withdrawn=False,
                )
            else:
                persist_claimed()
        finally:
            self._store.release_feishu_task(task, lease_owner)
        return task

    def _deliver_terminal_text(
        self,
        task: FeishuAgentTask,
        text: str,
        progress_sink: AgentProgressSink,
        persist: Callable[[], None],
        *,
        source_withdrawn: bool,
    ) -> bool:
        """Queue and attempt the terminal answer without changing its content."""
        return self._deliver_text_record(
            task,
            delivery_id="terminal-answer",
            text=text,
            persisted_text=task.answer or text,
            render_mode=(
                "terminal_with_visualization"
                if task.visualization is not None and self._public_app_url
                else "plain"
            ),
            progress_sink=progress_sink,
            persist=persist,
            source_withdrawn=source_withdrawn,
        )

    def _queue_terminal_text(
        self,
        task: FeishuAgentTask,
        text: str,
    ) -> None:
        """Queue the terminal answer before the task's successful state is stored."""
        self._queue_text_record(
            task,
            delivery_id="terminal-answer",
            text=task.answer or text,
            render_mode=(
                "terminal_with_visualization"
                if task.visualization is not None and self._public_app_url
                else "plain"
            ),
        )

    def _queue_text_record(
        self,
        task: FeishuAgentTask,
        *,
        delivery_id: str,
        text: str,
        render_mode: Literal["plain", "terminal_with_visualization"],
    ) -> tuple[FeishuDeliveryRecord, bool]:
        """Return the stable text intent and whether it was newly inserted."""
        target_message_id = task.progress_message_id or task.request.source_message_id
        kind: Literal["text_reply", "text_update"] = (
            "text_update" if task.progress_message_id else "text_reply"
        )
        content_sha256 = hashlib.sha256(
            f"{render_mode}\0{text}".encode("utf-8")
        ).hexdigest()
        record = self._delivery_record(task, delivery_id)
        if record is not None:
            if record.content_sha256 != content_sha256:
                raise ValueError(
                    f"Delivery identifier changed content: {delivery_id}"
                )
            return record, False
        now = datetime.now(timezone.utc)
        record = FeishuDeliveryRecord(
            delivery_id=delivery_id,
            kind=kind,
            status="pending",
            target_message_id=target_message_id,
            content_sha256=content_sha256,
            text=text,
            render_mode=render_mode,
            created_at=now,
            updated_at=now,
        )
        self._replace_delivery_record(task, record)
        return record, True

    def _deliver_text_record(
        self,
        task: FeishuAgentTask,
        *,
        delivery_id: str,
        text: str,
        persisted_text: Optional[str] = None,
        render_mode: Literal["plain", "terminal_with_visualization"] = "plain",
        progress_sink: AgentProgressSink,
        persist: Callable[[], None],
        source_withdrawn: bool,
    ) -> bool:
        """Persist a text intent before applying one reply or idempotent update."""
        semantic_text = persisted_text if persisted_text is not None else text
        record, created = self._queue_text_record(
            task,
            delivery_id=delivery_id,
            text=semantic_text,
            render_mode=render_mode,
        )
        if record.status in {"sent", "skipped"}:
            return record.status == "sent"
        if created:
            persist()
        if source_withdrawn:
            record.status = "skipped"
            record.updated_at = datetime.now(timezone.utc)
            record.last_error = "Source message was withdrawn."
            self._replace_delivery_record(task, record)
            persist()
            return False
        record.status = "pending"
        record.attempt_count += 1
        record.updated_at = datetime.now(timezone.utc)
        self._replace_delivery_record(task, record)
        persist()
        try:
            if record.kind == "text_update":
                progress_sink.update(record.target_message_id, text)
            else:
                record.remote_message_id = progress_sink.reply(
                    record.target_message_id,
                    text,
                )
        except FeishuSourceMessageWithdrawnError:
            record.status = "skipped"
            record.last_error = "Source message was withdrawn."
        except Exception as exc:
            logger.exception(
                "feishu_agent_terminal_text_delivery_failed task_id=%s delivery_id=%s",
                task.task_id,
                delivery_id,
            )
            record.status = "failed"
            record.last_error = str(exc)[:500]
        else:
            record.status = "sent"
            record.last_error = None
            record.delivered_at = datetime.now(timezone.utc)
        record.updated_at = datetime.now(timezone.utc)
        self._replace_delivery_record(task, record)
        persist()
        return record.status == "sent"

    def _deliver_terminal_artifact(
        self,
        task: FeishuAgentTask,
        progress_sink: AgentProgressSink,
        persist: Callable[[], None],
        *,
        artifact_path: Optional[Path] = None,
        source_withdrawn: bool,
    ) -> bool:
        """Persist and attempt one workbook delivery using archived bytes on retry."""
        artifact_name = task.artifact_name
        if artifact_name is None:
            return True
        record = self._delivery_record(task, "terminal-artifact")
        artifact_bytes = (
            artifact_path.read_bytes()
            if artifact_path is not None and artifact_path.exists()
            else self._store.get_artifact(task.task_id, artifact_name)
        )
        if artifact_bytes is None:
            logger.error(
                "feishu_agent_terminal_artifact_missing task_id=%s artifact=%s",
                task.task_id,
                artifact_name,
            )
            if record is not None:
                record.status = "failed"
                record.attempt_count = MAX_TERMINAL_DELIVERY_ATTEMPTS
                record.last_error = "Persisted terminal artifact is unavailable."
                record.updated_at = datetime.now(timezone.utc)
                self._replace_delivery_record(task, record)
                persist()
            return False
        record, created = self._queue_artifact_record(task, artifact_bytes)
        if record.status in {"sent", "skipped"}:
            return record.status == "sent"
        if created:
            persist()
        if source_withdrawn:
            record.status = "skipped"
            record.updated_at = datetime.now(timezone.utc)
            record.last_error = "Source message was withdrawn."
            self._replace_delivery_record(task, record)
            persist()
            return False
        record.status = "pending"
        record.attempt_count += 1
        record.updated_at = datetime.now(timezone.utc)
        self._replace_delivery_record(task, record)
        persist()
        try:
            if artifact_path is not None and artifact_path.exists():
                progress_sink.reply_file(record.target_message_id, artifact_path)
            else:
                with tempfile.TemporaryDirectory(prefix="feishu-delivery-") as temp_dir:
                    retry_path = Path(temp_dir) / artifact_name
                    retry_path.write_bytes(artifact_bytes)
                    progress_sink.reply_file(record.target_message_id, retry_path)
        except FeishuSourceMessageWithdrawnError:
            record.status = "skipped"
            record.last_error = "Source message was withdrawn."
        except Exception as exc:
            logger.exception(
                "feishu_agent_terminal_artifact_delivery_failed task_id=%s artifact=%s",
                task.task_id,
                artifact_name,
            )
            record.status = "failed"
            record.last_error = str(exc)[:500]
        else:
            record.status = "sent"
            record.last_error = None
            record.delivered_at = datetime.now(timezone.utc)
        record.updated_at = datetime.now(timezone.utc)
        self._replace_delivery_record(task, record)
        persist()
        return record.status == "sent"

    def _queue_artifact_record(
        self,
        task: FeishuAgentTask,
        artifact_bytes: bytes,
    ) -> tuple[FeishuDeliveryRecord, bool]:
        """Return the immutable artifact intent and whether it was newly inserted."""
        if task.artifact_name is None:
            raise ValueError("Terminal artifact name is required before queueing.")
        content_sha256 = hashlib.sha256(artifact_bytes).hexdigest()
        record = self._delivery_record(task, "terminal-artifact")
        if record is not None:
            if record.content_sha256 != content_sha256:
                raise ValueError("Terminal artifact changed after it was queued.")
            return record, False
        now = datetime.now(timezone.utc)
        record = FeishuDeliveryRecord(
            delivery_id="terminal-artifact",
            kind="artifact",
            status="pending",
            target_message_id=task.request.source_message_id,
            content_sha256=content_sha256,
            artifact_name=task.artifact_name,
            created_at=now,
            updated_at=now,
        )
        self._replace_delivery_record(task, record)
        return record, True

    def _resume_terminal_deliveries(
        self,
        task: FeishuAgentTask,
        progress_sink: AgentProgressSink,
        persist: Callable[[], None],
    ) -> FeishuAgentTask:
        """Retry unsent terminal outbox records without rerunning research."""
        retryable = [
            record
            for record in task.delivery_outbox
            if record.status in {"pending", "failed"}
            and record.attempt_count < MAX_TERMINAL_DELIVERY_ATTEMPTS
        ]
        for record in retryable:
            if record.kind == "artifact":
                delivered = self._deliver_terminal_artifact(
                    task,
                    progress_sink,
                    persist,
                    source_withdrawn=False,
                )
                if delivered and task.progress_message == "研究完成，但附件发送失败。":
                    task.progress_message = "研究完成。"
                    persist()
                    if task.answer and task.progress_message_id:
                        self._deliver_text_record(
                            task,
                            delivery_id="artifact-delivery-recovered",
                            text=task.answer + "\n\n附件已自动补发成功。",
                            progress_sink=progress_sink,
                            persist=persist,
                            source_withdrawn=False,
                        )
            elif record.text is not None:
                retry_text = record.text
                if record.render_mode == "terminal_with_visualization":
                    token = secrets.token_urlsafe(32)
                    task.visualization_token_hash = research_visualization_token_hash(
                        token
                    )
                    task.visualization_expires_at = (
                        datetime.now(timezone.utc)
                        + RESEARCH_VISUALIZATION_LINK_LIFETIME
                    )
                    query = urlencode({"token": token})
                    retry_text += (
                        "\n\n研究结果页面（30天内有效，点击后直接查看）：\n"
                        f"{self._public_app_url}/research/{task.task_id}?{query}"
                    )
                    persist()
                self._deliver_text_record(
                    task,
                    delivery_id=record.delivery_id,
                    text=retry_text,
                    persisted_text=record.text,
                    render_mode=record.render_mode,
                    progress_sink=progress_sink,
                    persist=persist,
                    source_withdrawn=False,
                )
        return task

    @staticmethod
    def _delivery_record(
        task: FeishuAgentTask,
        delivery_id: str,
    ) -> Optional[FeishuDeliveryRecord]:
        """Return one isolated delivery record from the task outbox."""
        return next(
            (
                record.model_copy(deep=True)
                for record in task.delivery_outbox
                if record.delivery_id == delivery_id
            ),
            None,
        )

    @staticmethod
    def _replace_delivery_record(
        task: FeishuAgentTask,
        record: FeishuDeliveryRecord,
    ) -> None:
        """Insert or replace one task-local delivery intent by stable identifier."""
        task.delivery_outbox = [
            existing
            for existing in task.delivery_outbox
            if existing.delivery_id != record.delivery_id
        ] + [record]


class FeishuAgentOutcome(BaseModel):
    """Terminal conversational answer and optional generated workbook."""

    model_config = ConfigDict(arbitrary_types_allowed=True, extra="forbid")

    answer: str = Field(min_length=1)
    artifact_path: Optional[Path] = None
    visualization: Optional[FeishuResearchVisualization] = None


def research_visualization_token_hash(token: str) -> str:
    """Return the stable digest used to verify one viewer bearer token."""
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def has_retryable_terminal_deliveries(task: FeishuAgentTask) -> bool:
    """Return whether a successful task needs one bounded worker retry."""
    return task.status == AnalysisTaskStatus.SUCCEEDED and any(
        record.status in {"pending", "failed"}
        and record.attempt_count < MAX_TERMINAL_DELIVERY_ATTEMPTS
        for record in task.delivery_outbox
    )


class ResearchToolbox:
    """Own audited datasets and deterministic transformations for one agent turn."""

    def __init__(
        self,
        provider: MarketDataProvider,
        request_id: str,
        *,
        python_sandbox: Optional[PythonSandbox] = None,
        artifact_dir: Optional[Path] = None,
        dataset_archive: Optional[AgentTaskStore] = None,
        task_id: str = "",
        session_dataset: Optional[QueryResult] = None,
    ) -> None:
        """Store provider access and the optional secretless Python boundary."""
        self._provider = provider
        self._request_id = request_id
        self._python_sandbox = python_sandbox
        self._artifact_dir = artifact_dir
        self._dataset_archive = dataset_archive
        self._task_id = task_id
        self._datasets: Dict[str, QueryResult] = {}
        self._manifest: Optional[ResearchManifest] = None
        if self._dataset_archive is not None and self._task_id:
            self._manifest = self._dataset_archive.get_research_manifest(self._task_id)
            if self._manifest is None:
                self._manifest = new_research_manifest(self._task_id)
                self._dataset_archive.put_research_manifest(self._manifest)
        if session_dataset is not None:
            retained_session = session_dataset.model_copy(
                update={"query_id": "session_dataset"},
                deep=True,
            )
            self._datasets["session_dataset"] = retained_session
            self._record_dataset_evidence(
                retained_session,
                source_dataset_ids=[],
                operation_parameters={"source": "session_workspace"},
                requested_fields=list(retained_session.columns),
            )

    @property
    def definitions(self) -> List[Dict[str, Any]]:
        """Return provider-neutral research tools advertised to the model."""
        pipeline_schema = ResultPipeline.model_json_schema()
        pipeline_definitions = pipeline_schema.pop("$defs", {})
        transform_parameters: Dict[str, Any] = {
            "type": "object",
            "properties": {"pipeline": pipeline_schema},
            "required": ["pipeline"],
            "additionalProperties": False,
        }
        if pipeline_definitions:
            # JSON Schema references resolve from the tool-input root, so move
            # Pydantic's definitions beside the wrapper object.
            transform_parameters["$defs"] = pipeline_definitions
        definitions = [
            {
                "type": "function",
                "function": {
                    "name": "request_clarification",
                    "description": (
                        "Ask the user to resolve material ambiguity before any data "
                        "query. Provide two to four concrete choices and mark the "
                        "safest default as recommended."
                    ),
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "question": {"type": "string"},
                            "options": {
                                "type": "array",
                                "items": {"type": "string"},
                                "minItems": 2,
                                "maxItems": 4,
                            },
                        },
                        "required": ["question", "options"],
                        "additionalProperties": False,
                    },
                },
            },
            {
                "type": "function",
                "function": {
                    "name": "search_market_data",
                    "description": (
                        "Find relevant read-only operations from the configured "
                        "market-data provider before querying."
                    ),
                    "parameters": {
                        "type": "object",
                        "properties": {"query": {"type": "string"}},
                        "required": ["query"],
                        "additionalProperties": False,
                    },
                },
            },
            {
                "type": "function",
                "function": {
                    "name": "query_market_data",
                    "description": (
                        "Execute one provider-supported read-only operation and retain "
                        "the complete returned dataset for later calculations."
                    ),
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "operation": {"type": "string"},
                            "params": {"type": "object"},
                            "fields": {"type": "array", "items": {"type": "string"}},
                        },
                        "required": ["operation", "params", "fields"],
                        "additionalProperties": False,
                    },
                },
            },
            {
                "type": "function",
                "function": {
                    "name": "rank_dataset",
                    "description": "Sort a retained dataset, keep a bounded top or bottom set, and optionally select output fields.",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "dataset_id": {"type": "string"},
                            "sort_by": {"type": "string"},
                            "direction": {"type": "string", "enum": ["asc", "desc"]},
                            "limit": {"type": "integer", "minimum": 1, "maximum": 1000},
                            "fields": {"type": "array", "items": {"type": "string"}},
                        },
                        "required": ["dataset_id", "sort_by", "direction", "limit", "fields"],
                        "additionalProperties": False,
                    },
                },
            },
            {
                "type": "function",
                "function": {
                    "name": "join_datasets",
                    "description": "Join two retained datasets on explicit key fields using validated many-to-one or one-to-one semantics.",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "left_dataset_id": {"type": "string"},
                            "right_dataset_id": {"type": "string"},
                            "join_on": {"type": "array", "items": {"type": "string"}},
                            "right_fields": {"type": "array", "items": {"type": "string"}},
                            "cardinality": {"type": "string", "enum": ["one_to_one", "many_to_one"]},
                        },
                        "required": ["left_dataset_id", "right_dataset_id", "join_on", "right_fields", "cardinality"],
                        "additionalProperties": False,
                    },
                },
            },
            {
                "type": "function",
                "function": {
                    "name": "transform_dataset",
                    "description": (
                        "Apply an allowlisted ResultPipeline to retained datasets. "
                        "Use source_query_id and output_query_id plus 1-16 validated steps."
                    ),
                    "parameters": transform_parameters,
                },
            },
            {
                "type": "function",
                "function": {
                    "name": "export_excel",
                    "description": (
                        "Create a polished Excel workbook from a retained "
                        "dataset. methodology must enumerate, for every derived "
                        "indicator: its input fields, the exact price series and "
                        "adjustment basis, window semantics (trading days or "
                        "calendar days), and the event-deduplication rule. Never "
                        "refer to another indicator's basis with phrases like "
                        "same basis; restate the full computation each time. "
                        "column_notes must explain every exported column in "
                        "concise Chinese so a reader can interpret each value."
                    ),
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "dataset_id": {"type": "string"},
                            "title": {"type": "string"},
                            "methodology": {"type": "string"},
                            "column_notes": {
                                "type": "array",
                                "items": {
                                    "type": "object",
                                    "properties": {
                                        "column": {"type": "string"},
                                        "note": {"type": "string"},
                                    },
                                    "required": ["column", "note"],
                                    "additionalProperties": False,
                                },
                                "minItems": 1,
                                "maxItems": 200,
                            },
                        },
                        "required": [
                            "dataset_id",
                            "title",
                            "methodology",
                            "column_notes",
                        ],
                        "additionalProperties": False,
                    },
                },
            },
        ]
        if "session_dataset" in self._datasets:
            definitions.insert(
                0,
                {
                    "type": "function",
                    "function": {
                        "name": "inspect_session_dataset",
                        "description": (
                            "Inspect the complete final dataset retained from the "
                            "previous successful turn in this named session. Use "
                            "dataset_id session_dataset for follow-up filtering, "
                            "ranking, joining, or Python analysis."
                        ),
                        "parameters": {
                            "type": "object",
                            "properties": {},
                            "additionalProperties": False,
                        },
                    },
                },
            )
        if self._python_sandbox is not None:
            definitions.insert(
                -1,
                {
                    "type": "function",
                    "function": {
                        "name": "run_python_analysis",
                        "description": (
                            "Run a restricted pandas/numpy DataFrame program in the "
                            "independent secretless sandbox. Inputs are available as "
                            "datasets[dataset_id], imports and external I/O are forbidden, "
                            "and the final DataFrame must be assigned to result. The "
                            "preloaded helper research_flag_overlapping_signals(frame, "
                            "ts_code, signal_date, window_end) returns the Boolean "
                            "keep-mask that collapses overlapping same-security signal "
                            "windows into one keep-first event."
                        ),
                        "parameters": {
                            "type": "object",
                            "properties": {
                                "dataset_ids": {
                                    "type": "array",
                                    "items": {"type": "string"},
                                    "minItems": 1,
                                    "maxItems": 8,
                                },
                                "code": {
                                    "type": "string",
                                    "maxLength": 12_000,
                                },
                            },
                            "required": ["dataset_ids", "code"],
                            "additionalProperties": False,
                        },
                    },
                },
            )
        return definitions

    def call(
        self,
        name: str,
        arguments: Mapping[str, Any],
        progress: Callable[[str, str], None],
    ) -> Dict[str, Any]:
        """Execute one named tool or fail fast on an unknown capability."""
        if name == "inspect_session_dataset":
            result = self._datasets.get("session_dataset")
            if result is None:
                raise ValueError("This named session has no retained dataset.")
            return _result_payload(result)
        if name == "request_clarification":
            return {
                "clarification": _format_clarification(arguments),
                "instruction": (
                    "Return the clarification verbatim as the final answer and do "
                    "not call another tool in this turn."
                ),
            }
        if name == "search_market_data":
            operations = self._provider.search_operations(str(arguments["query"]))
            operation_results = []
            for operation in operations:
                query_shapes = self._provider.describe_query_shapes(operation.name)
                if not query_shapes:
                    continue
                operation_results.append(
                    {
                        "name": operation.name,
                        "description": operation.description,
                        "query_shapes": list(query_shapes),
                    }
                )
            return {
                "operations": operation_results
            }
        if name == "query_market_data":
            operation = str(arguments["operation"])
            if not self._provider.supports(operation):
                raise ValueError(f"Unsupported market-data operation: {operation}")
            params = dict(arguments["params"])
            fields = [str(field) for field in arguments["fields"]]
            self._provider.validate_query(operation, params, fields)
            progress("querying", f"正在查询市场数据：{operation}…")
            dataset_id = f"dataset_{len(self._datasets) + 1}"
            frame = self._provider.query(
                operation,
                params,
                fields,
                api_route="/feishu/agent/tools/query",
                request_id=self._request_id,
                query_id=dataset_id,
            )
            result = _frame_to_result(
                dataset_id,
                str(frame.attrs.get("provider") or self._provider.name),
                operation,
                frame,
            )
            self._retain_dataset(
                result,
                source_dataset_ids=[],
                operation_parameters=params,
                requested_fields=fields,
            )
            return _result_payload(result)
        if name == "rank_dataset":
            progress("calculating", "正在执行排序与排名…")
            dataset_id = str(arguments["dataset_id"])
            source = self._datasets.get(dataset_id)
            if source is None:
                raise ValueError(f"Unknown ranking dataset: {dataset_id}")
            sort_by = str(arguments["sort_by"])
            fields = [str(field) for field in arguments["fields"]]
            missing = set([sort_by] + fields).difference(source.columns)
            if missing:
                raise ValueError(
                    "Ranking fields are missing: " + ", ".join(sorted(missing))
                )
            frame = pd.DataFrame(source.rows)
            ranked = frame.sort_values(
                sort_by,
                ascending=str(arguments["direction"]) == "asc",
                kind="mergesort",
                na_position="last",
            ).head(int(arguments["limit"]))
            if fields:
                ranked = ranked.loc[:, fields]
            output_id = f"dataset_{len(self._datasets) + 1}"
            result = _frame_to_result(
                output_id,
                source.provider,
                f"{source.operation}_ranked",
                ranked,
            )
            self._retain_dataset(
                result,
                source_dataset_ids=[dataset_id],
                operation_parameters={
                    "sort_by": sort_by,
                    "direction": str(arguments["direction"]),
                    "limit": int(arguments["limit"]),
                    "fields": fields,
                },
                requested_fields=fields,
            )
            return _result_payload(result)
        if name == "join_datasets":
            progress("calculating", "正在合并市场与财务数据…")
            left = self._datasets.get(str(arguments["left_dataset_id"]))
            right = self._datasets.get(str(arguments["right_dataset_id"]))
            if left is None or right is None:
                raise ValueError("Join references an unknown dataset.")
            join_on = [str(field) for field in arguments["join_on"]]
            right_fields = [str(field) for field in arguments["right_fields"]]
            missing_left = set(join_on).difference(left.columns)
            missing_right = set(join_on + right_fields).difference(right.columns)
            if missing_left or missing_right:
                raise ValueError("Join fields are missing from the retained datasets.")
            cardinality = str(arguments["cardinality"])
            validate = "1:1" if cardinality == "one_to_one" else "m:1"
            merged = pd.DataFrame(left.rows).merge(
                pd.DataFrame(right.rows).loc[:, join_on + right_fields],
                how="inner",
                on=join_on,
                validate=validate,
            )
            output_id = f"dataset_{len(self._datasets) + 1}"
            result = _frame_to_result(
                output_id,
                left.provider,
                f"{left.operation}_joined",
                merged,
            )
            self._retain_dataset(
                result,
                source_dataset_ids=[left.query_id, right.query_id],
                operation_parameters={
                    "join_on": join_on,
                    "right_fields": right_fields,
                    "cardinality": cardinality,
                },
                requested_fields=list(result.columns),
            )
            return _result_payload(result)
        if name == "transform_dataset":
            progress("calculating", "正在执行确定性筛选与计算…")
            pipeline = ResultPipeline.model_validate(arguments["pipeline"])
            source = self._datasets.get(pipeline.source_query_id)
            if source is None:
                raise ValueError(
                    f"Unknown source dataset: {pipeline.source_query_id}"
                )
            result = ResultPipelineExecutor().execute(
                pipeline,
                source,
                self._datasets,
            )
            pipeline_sources = [pipeline.source_query_id]
            for step in pipeline.steps:
                right_source = step.right_source_query_id
                if right_source and right_source not in pipeline_sources:
                    pipeline_sources.append(right_source)
            self._retain_dataset(
                result,
                source_dataset_ids=pipeline_sources,
                operation_parameters={"pipeline": pipeline.model_dump(mode="json")},
                requested_fields=list(result.columns),
            )
            return _result_payload(result)
        if name == "run_python_analysis":
            if self._python_sandbox is None:
                raise ValueError("The independent Python sandbox is not configured.")
            dataset_ids = [str(value) for value in arguments["dataset_ids"]]
            if len(dataset_ids) != len(set(dataset_ids)):
                raise ValueError("Python sandbox dataset identifiers must be unique.")
            missing = [
                dataset_id
                for dataset_id in dataset_ids
                if dataset_id not in self._datasets
            ]
            if missing:
                raise ValueError(
                    "Python sandbox references unknown datasets: "
                    + ", ".join(missing)
                )
            progress("calculating", "正在安全沙箱中执行通用数据计算…")
            output_id = f"dataset_{len(self._datasets) + 1}"
            frames = {
                dataset_id: pd.DataFrame(self._datasets[dataset_id].rows)
                for dataset_id in dataset_ids
            }
            frame = self._python_sandbox.run(str(arguments["code"]), frames)
            result = _frame_to_result(
                output_id,
                "research_sandbox",
                "python_dataframe",
                frame,
            )
            self._retain_dataset(
                result,
                source_dataset_ids=dataset_ids,
                operation_parameters={
                    "code_sha256": hashlib.sha256(
                        str(arguments["code"]).encode("utf-8")
                    ).hexdigest()
                },
                requested_fields=list(result.columns),
            )
            return _result_payload(result)
        if name == "export_excel":
            dataset_id = str(arguments["dataset_id"])
            result = self._datasets.get(dataset_id)
            if result is None:
                raise ValueError(f"Unknown Excel dataset: {dataset_id}")
            column_notes = _validated_column_notes(arguments["column_notes"], result)
            progress("exporting", "正在生成 Excel 研究结果…")
            artifact_path = build_research_workbook(
                result,
                str(arguments["title"]),
                str(arguments["methodology"]),
                output_dir=self._artifact_dir,
                column_notes=column_notes,
            )
            if self._dataset_archive is not None and self._task_id:
                if dataset_id == "session_dataset":
                    self._dataset_archive.archive_dataset(self._task_id, result)
                self._dataset_archive.mark_final_dataset(self._task_id, dataset_id)
                if self._manifest is None:
                    self._manifest = new_research_manifest(self._task_id)
                self._manifest = select_final_dataset(self._manifest, dataset_id)
                self._dataset_archive.put_research_manifest(self._manifest)
            return {
                "file_name": artifact_path.name,
                "file_path": str(artifact_path),
                "row_count": result.row_count,
            }
        raise ValueError(f"Unknown research tool: {name}")

    def _retain_dataset(
        self,
        result: QueryResult,
        *,
        source_dataset_ids: List[str],
        operation_parameters: Mapping[str, Any],
        requested_fields: List[str],
    ) -> None:
        """Retain a complete dataset in memory and the durable task archive."""
        self._datasets[result.query_id] = result
        if self._dataset_archive is not None and self._task_id:
            self._dataset_archive.archive_dataset(self._task_id, result)
        self._record_dataset_evidence(
            result,
            source_dataset_ids=source_dataset_ids,
            operation_parameters=operation_parameters,
            requested_fields=requested_fields,
        )

    def _record_dataset_evidence(
        self,
        result: QueryResult,
        *,
        source_dataset_ids: List[str],
        operation_parameters: Mapping[str, Any],
        requested_fields: List[str],
    ) -> None:
        """Persist executor-derived evidence independently of model narration."""
        if self._dataset_archive is None or not self._task_id:
            return
        if self._manifest is None:
            self._manifest = new_research_manifest(self._task_id)
        evidence = research_dataset_evidence(
            result,
            source_dataset_ids=source_dataset_ids,
            operation_parameters=operation_parameters,
            requested_fields=requested_fields,
        )
        self._manifest = record_dataset_evidence(self._manifest, evidence)
        self._dataset_archive.put_research_manifest(self._manifest)


def _validated_column_notes(
    raw_notes: object,
    result: QueryResult,
) -> Dict[str, str]:
    """Return one bounded note per column and reject incomplete coverage."""
    if not isinstance(raw_notes, list) or not raw_notes:
        raise ValueError("export_excel requires column_notes for every column.")
    notes: Dict[str, str] = {}
    for entry in raw_notes:
        if not isinstance(entry, Mapping):
            raise ValueError("Each column note must be an object with column and note.")
        column = str(entry.get("column") or "").strip()
        note = str(entry.get("note") or "").strip()
        if not column:
            raise ValueError("Each column note requires a column name.")
        if not MIN_COLUMN_NOTE_CHARACTERS <= len(note) <= MAX_COLUMN_NOTE_CHARACTERS:
            raise ValueError(
                f"Column note for {column} must be "
                f"{MIN_COLUMN_NOTE_CHARACTERS}-{MAX_COLUMN_NOTE_CHARACTERS} characters."
            )
        if column in notes:
            raise ValueError(f"Duplicate column note for {column}.")
        notes[column] = note
    missing = sorted(set(result.columns).difference(notes))
    unexpected = sorted(set(notes).difference(result.columns))
    if missing or unexpected:
        parts = []
        if missing:
            parts.append("missing notes for: " + ", ".join(missing))
        if unexpected:
            parts.append("notes for unknown columns: " + ", ".join(unexpected))
        raise ValueError(
            "column_notes must exactly cover the exported columns ("
            + "; ".join(parts)
            + ")."
        )
    return notes


RESEARCH_COLUMN_NOTES_SHEET_NAME = "列说明"

_DATE_COLUMN_TOKEN_PATTERN = re.compile(r"date|day|time|时|日|期", re.IGNORECASE)
_COMPACT_CALENDAR_DATE_PATTERN = re.compile(r"\d{8}")


def compact_calendar_date_text(column: str, value: Any) -> Any:
    """Render an eight-digit YYYYMMDD calendar value as YYYY-MM-DD.

    Research datasets frequently carry calendar dates as plain integers such
    as 20260917; numeric renderings split those into thousands groups
    (20,260,917). Any date-named column whose value is a valid calendar date
    keeps its calendar shape, and every other value passes through unchanged.
    """
    if not _DATE_COLUMN_TOKEN_PATTERN.search(str(column or "")):
        return value
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, str):
        text = value.strip()
    else:
        return value
    if not _COMPACT_CALENDAR_DATE_PATTERN.fullmatch(text):
        return value
    try:
        parsed = datetime.strptime(text, "%Y%m%d")
    except ValueError:
        return value
    return parsed.strftime("%Y-%m-%d")


def build_research_workbook(
    result: QueryResult,
    title: str,
    methodology: str,
    *,
    output_dir: Optional[Path] = None,
    column_notes: Optional[Mapping[str, str]] = None,
) -> Path:
    """Create one readable workbook with results, notes, and methodology."""
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Excel export requires the openpyxl dependency.") from exc

    notes = {str(key): str(value) for key, value in (column_notes or {}).items()}
    workbook = Workbook()
    results_sheet = workbook.active
    results_sheet.title = "Results"
    methodology_sheet = workbook.create_sheet("Methodology")
    notes_sheet = workbook.create_sheet(RESEARCH_COLUMN_NOTES_SHEET_NAME)
    results_sheet.sheet_view.showGridLines = False
    methodology_sheet.sheet_view.showGridLines = False
    notes_sheet.sheet_view.showGridLines = False

    results_sheet["A2"] = title
    results_sheet["A2"].font = Font(name="Arial", size=14, bold=True)
    results_sheet["A3"] = f"Rows: {result.row_count}"
    results_sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="666666")
    header_row = 5
    for column_index, column in enumerate(result.columns, start=1):
        cell = results_sheet.cell(header_row, column_index, column)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        note = notes.get(column)
        if note:
            cell.comment = Comment(note, "A股研究助手", height=120, width=280)
    linkable_columns = {
        column
        for column in result.columns
        if is_security_code_column(column, [row.get(column) for row in result.rows[:50]])
    }
    for row_index, row in enumerate(result.rows, start=header_row + 1):
        for column_index, column in enumerate(result.columns, start=1):
            cell = results_sheet.cell(
                row_index,
                column_index,
                compact_calendar_date_text(column, row.get(column)),
            )
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
            if column in linkable_columns:
                quote_url = security_quote_page_url(row.get(column))
                if quote_url:
                    cell.hyperlink = quote_url
                    cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
    results_sheet.freeze_panes = "A6"
    results_sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(max(len(result.columns), 1))}"
        f"{header_row + max(result.row_count, 1)}"
    )
    for column_index, column in enumerate(result.columns, start=1):
        sample_values = [str(row.get(column) or "") for row in result.rows[:200]]
        width = min(max([len(column)] + [len(value) for value in sample_values]) + 2, 36)
        results_sheet.column_dimensions[get_column_letter(column_index)].width = width
        normalized_column = column.casefold()
        if "ratio" in normalized_column or "yield" in normalized_column:
            for row_index in range(header_row + 1, header_row + result.row_count + 1):
                results_sheet.cell(row_index, column_index).number_format = "0.00%"
    results_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    results_sheet.page_setup.orientation = "landscape"
    results_sheet.page_setup.fitToWidth = 1
    results_sheet.page_setup.fitToHeight = 0
    results_sheet.print_area = (
        f"A1:{get_column_letter(max(len(result.columns), 1))}"
        f"{header_row + max(result.row_count, 1)}"
    )

    notes_sheet["A2"] = "列说明"
    notes_sheet["A2"].font = Font(name="Arial", size=14, bold=True)
    notes_header_row = 4
    for column_index, header in enumerate(("列", "说明"), start=1):
        cell = notes_sheet.cell(notes_header_row, column_index, header)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, column in enumerate(result.columns, start=notes_header_row + 1):
        note_cell = notes_sheet.cell(row_index, 1, column)
        note_cell.font = Font(name="Arial", size=10, bold=True)
        note_cell.alignment = Alignment(vertical="center")
        detail_cell = notes_sheet.cell(row_index, 2, notes.get(column, ""))
        detail_cell.font = Font(name="Arial", size=10)
        detail_cell.alignment = Alignment(wrap_text=True, vertical="top")
    notes_sheet.column_dimensions["A"].width = 24
    notes_sheet.column_dimensions["B"].width = 80
    notes_sheet.freeze_panes = "A5"
    notes_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    notes_sheet.page_setup.fitToWidth = 1
    notes_sheet.page_setup.fitToHeight = 0
    notes_sheet.print_area = (
        f"A1:B{notes_header_row + max(len(result.columns), 1)}"
    )

    methodology_sheet["A2"] = "Methodology"
    methodology_sheet["A2"].font = Font(name="Arial", size=14, bold=True)
    methodology_sheet["A4"] = "Data provider"
    methodology_sheet["B4"] = result.provider
    methodology_sheet["A5"] = "Operation"
    methodology_sheet["B5"] = result.operation
    methodology_sheet["A6"] = "Dataset"
    methodology_sheet["B6"] = result.query_id
    methodology_sheet["A7"] = "Method"
    methodology_sheet["B7"] = methodology
    methodology_sheet["A8"] = "Generated at"
    methodology_sheet["B8"] = datetime.now(timezone.utc).replace(tzinfo=None)
    methodology_sheet["B8"].number_format = "yyyy-mm-dd hh:mm"
    methodology_sheet.column_dimensions["A"].width = 18
    methodology_sheet.column_dimensions["B"].width = 60
    methodology_sheet["B7"].alignment = Alignment(wrap_text=True, vertical="top")
    for row in methodology_sheet.iter_rows(min_row=4, max_row=8, min_col=1, max_col=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
    for row_index in range(4, 9):
        methodology_sheet.cell(row_index, 1).font = Font(name="Arial", size=10, bold=True)
    methodology_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    methodology_sheet.page_setup.fitToWidth = 1
    methodology_sheet.page_setup.fitToHeight = 1
    methodology_sheet.print_area = "A1:B9"

    target_dir = output_dir or Path(tempfile.mkdtemp(prefix="feishu-agent-"))
    target_dir.mkdir(parents=True, exist_ok=True)
    output_path = target_dir / "a_share_research.xlsx"
    workbook.save(output_path)
    return output_path


def _frame_to_result(
    query_id: str,
    provider: str,
    operation: str,
    frame: pd.DataFrame,
) -> QueryResult:
    """Convert one provider frame into the shared deterministic result contract."""
    normalized = frame.astype(object).where(pd.notna(frame), None)
    rows = normalized.to_dict(orient="records")
    return QueryResult(
        query_id=query_id,
        provider=provider,
        operation=operation,
        status=QueryStatus.SUCCESS,
        columns=list(frame.columns),
        rows=rows,
        row_count=len(rows),
        completeness="complete",
        completeness_evidence=["feishu_agent_complete_provider_query"],
    )


def _result_payload(
    result: QueryResult,
    as_of: Optional[str] = None,
) -> Dict[str, Any]:
    """Return bounded model-visible evidence while retaining the complete dataset."""
    payload = {
        "dataset_id": result.query_id,
        "row_count": result.row_count,
        "columns": result.columns,
        "summary": result.summary,
        "preview": result.rows[:MAX_AGENT_PREVIEW_ROWS],
        "preview_truncated": result.row_count > MAX_AGENT_PREVIEW_ROWS,
        "dataset_scope": "complete_retained_result",
        "preview_note": "Display only; tools use every row retained by dataset_id.",
    }
    # Unit facts travel with the data so every model sees them without any
    # runtime-specific prompt logic.
    field_units = field_unit_notes_for(result.operation, result.columns)
    if field_units:
        payload["field_units"] = field_units
    # Freshness travels the same way: the newest disclosure date in the dataset
    # makes stale event-triggered records visible before they are quoted as
    # current. Tests pin ``as_of`` explicitly; production uses Beijing's date.
    recency = data_recency_note(
        result.columns,
        result.rows,
        as_of or datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y%m%d"),
    )
    if recency is not None:
        payload["data_recency"] = recency
    return payload


def _format_clarification(arguments: Any) -> str:
    """Render one bounded clarification that can be answered by number or text."""
    if not isinstance(arguments, dict):
        raise RuntimeError("Research clarification arguments must be an object.")
    question = str(arguments.get("question") or "").strip()
    raw_options = arguments.get("options")
    if not question or not isinstance(raw_options, list):
        raise RuntimeError("Research clarification requires a question and options.")
    options = [str(option).strip() for option in raw_options if str(option).strip()]
    if not 2 <= len(options) <= 4 or len(options) != len(set(options)):
        raise RuntimeError(
            "Research clarification requires two to four unique options."
        )
    lines = [question]
    lines.extend(f"{index}. {option}" for index, option in enumerate(options, start=1))
    lines.append("请回复序号，或直接补充你的完整口径。")
    return "\n".join(lines)
