from datetime import datetime, timezone
from pathlib import Path
import os
import re

import pandas as pd
from openpyxl import load_workbook
import pytest

from china_a_share.bootstrap import create_feishu_agent_runtime
from china_a_share.capabilities import get_operation_capability
from china_a_share.config import Settings
from china_a_share.core.contracts import AnalysisTaskStatus, QueryResult, QueryStatus
from china_a_share.feishu_agent import (
    FeishuAgentCoordinator,
    FeishuAgentConversationTurn,
    FeishuAgentOutcome,
    FeishuAgentRequest,
    FeishuAgentTask,
    FeishuResearchVisualization,
    ResearchToolbox,
    build_research_workbook,
    research_visualization_token_hash,
)
from china_a_share.feishu import FeishuOpenApiClient
from china_a_share.registry import TushareOperationCatalog
from china_a_share.tasks import MemoryAnalysisTaskStore


class FakeResponse:
    def __init__(self, payload, status_code=200):
        self._payload = payload
        self.status_code = status_code
        self.text = ""

    def json(self):
        return self._payload


class SequenceSession:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    def post(self, url, **kwargs):
        self.calls.append((url, kwargs))
        return self.responses.pop(0)

    def put(self, url, **kwargs):
        self.calls.append((url, kwargs))
        return self.responses.pop(0)


class FakeOperation:
    name = "daily"
    description = "Daily A-share prices with trade_date and close fields."


class UnauditedFakeOperation:
    name = "bak_basic"
    description = "Unaudited backup operation."


class FakeProvider:
    name = "test-provider"

    def search_operations(self, _prompt):
        return [FakeOperation()]

    def supports(self, operation):
        return operation == "daily"

    def describe_query_shapes(self, operation):
        capability = get_operation_capability(operation)
        if capability is None:
            return ()
        return tuple(
            {
                "shape_id": shape.shape_id,
                "required_params": list(shape.required_params),
            }
            for shape in capability.query_shapes
        )

    def validate_query(self, operation, params, fields):
        assert operation == "daily"
        assert params == {"trade_date": "20260916"}
        assert fields == ["ts_code", "close"]

    def query(
        self,
        operation,
        params,
        fields,
        *,
        api_route,
        request_id,
        query_id,
    ):
        assert operation == "daily"
        assert params == {"trade_date": "20260916"}
        assert fields == ["ts_code", "close"]
        return pd.DataFrame(
            [
                {"ts_code": "000001.SZ", "close": 10.25},
                {"ts_code": "600000.SH", "close": 11.5},
            ]
        )


class RecentReturnProvider:
    name = "test-provider"

    def search_operations(self, prompt):
        return [FakeOperation()]

    def supports(self, operation):
        return operation in {"daily", "stock_basic", "trade_cal"}

    def validate_query(self, operation, params, fields):
        assert self.supports(operation)

    def query(
        self,
        operation,
        params,
        fields,
        *,
        api_route,
        request_id,
        query_id,
    ):
        if operation == "stock_basic":
            assert params == {"list_status": "L"}
            return pd.DataFrame(
                [
                    {"ts_code": "000001.SZ", "name": "Ping An Bank", "industry": "Bank"},
                    {"ts_code": "600000.SH", "name": "SPDB", "industry": "Bank"},
                ]
            )
        if operation == "trade_cal":
            assert params["exchange"] == "SSE"
            assert params["is_open"] == "1"
            return pd.DataFrame(
                [
                    {"cal_date": trade_date, "is_open": 1}
                    for trade_date in ("20260914", "20260915", "20260916")
                ]
            )
        assert operation == "daily"
        changes = {
            "20260914": (1.0, 3.0),
            "20260915": (2.0, -1.0),
            "20260916": (3.0, 1.0),
        }
        trade_dates = (
            [params["trade_date"]]
            if "trade_date" in params
            else list(changes)
        )
        rows = []
        for trade_date in trade_dates:
            first_change, second_change = changes[trade_date]
            rows.extend(
                [
                    {
                        "ts_code": "000001.SZ",
                        "trade_date": trade_date,
                        "close": 10.0,
                        "pct_chg": first_change,
                    },
                    {
                        "ts_code": "600000.SH",
                        "trade_date": trade_date,
                        "close": 12.0,
                        "pct_chg": second_change,
                    },
                ]
            )
        return pd.DataFrame(rows)


class RecordingPythonSandbox:
    def __init__(self):
        self.calls = []

    def run(self, code, datasets):
        self.calls.append((code, datasets))
        daily = datasets["dataset_1"].copy()
        daily["pct_chg"] = pd.to_numeric(daily["pct_chg"])
        selected_dates = sorted(daily["trade_date"].unique())[-3:]
        selected = daily.loc[daily["trade_date"].isin(selected_dates)]
        result = (
            selected.groupby("ts_code")["pct_chg"]
            .apply(lambda values: ((1 + values / 100).prod() - 1) * 100)
            .rename("period_return_pct")
            .reset_index()
            .sort_values("period_return_pct", ascending=False)
            .head(1)
        )
        return result


class RecordingDispatcher:
    def __init__(self):
        self.task_ids = []

    def dispatch(self, task_id):
        self.task_ids.append(task_id)


class RecordingSink:
    def __init__(self):
        self.messages = []
        self.updates = []
        self.files = []

    def reply(self, message_id, text):
        self.messages.append((message_id, text))
        return "progress-message-1"

    def update(self, message_id, text):
        self.updates.append((message_id, text))

    def reply_file(self, message_id, path):
        self.files.append((message_id, path))


def test_request_clarification_returns_one_bounded_prompt():
    toolbox = ResearchToolbox(FakeProvider(), "request-1")

    payload = toolbox.call(
        "request_clarification",
        {
            "question": "请确认排名口径：",
            "options": ["口径一（推荐）", "口径二", "自定义口径"],
        },
        lambda _stage, _message: None,
    )

    assert payload["clarification"] == (
        "请确认排名口径：\n"
        "1. 口径一（推荐）\n"
        "2. 口径二\n"
        "3. 自定义口径\n"
        "请回复序号，或直接补充你的完整口径。"
    )
    assert "final answer" in payload["instruction"]


def test_search_market_data_returns_all_audited_operations_with_query_shapes():
    class DiscoveryProvider(FakeProvider):
        def search_operations(self, prompt):
            assert prompt == "valuation"
            return [UnauditedFakeOperation(), FakeOperation()]

    payload = ResearchToolbox(DiscoveryProvider(), "request-1").call(
        "search_market_data",
        {"query": "valuation"},
        lambda _stage, _message: None,
    )

    assert [operation["name"] for operation in payload["operations"]] == ["daily"]
    assert payload["operations"][0]["query_shapes"] == [
        {
            "shape_id": "security",
            "required_params": ["ts_code"],
        },
        {
            "shape_id": "market_snapshot",
            "required_params": ["trade_date"],
        },
        {
            "shape_id": "bounded_range",
            "required_params": ["start_date", "end_date"],
        },
    ]


def test_search_market_data_exposes_production_catalog_operations():
    class CatalogProvider(FakeProvider):
        def search_operations(self, prompt):
            assert prompt == "valuation ranking"
            return TushareOperationCatalog().search(prompt)

    payload = ResearchToolbox(CatalogProvider(), "request-1").call(
        "search_market_data",
        {"query": "valuation ranking"},
        lambda _stage, _message: None,
    )

    operation_names = [operation["name"] for operation in payload["operations"]]
    assert "daily_basic" in operation_names
    assert "bak_basic" not in operation_names
    assert len(operation_names) > 12


def test_generic_query_and_python_sandbox_replace_prompt_specific_ranking_tool():
    progress = []
    sandbox = RecordingPythonSandbox()
    toolbox = ResearchToolbox(
        RecentReturnProvider(),
        "request-1",
        python_sandbox=sandbox,
    )

    daily = toolbox.call(
        "query_market_data",
        {
            "operation": "daily",
            "params": {"start_date": "20260901", "end_date": "20260916"},
            "fields": ["ts_code", "trade_date", "close", "pct_chg"],
        },
        lambda stage, message: progress.append((stage, message)),
    )
    payload = toolbox.call(
        "run_python_analysis",
        {
            "dataset_ids": [daily["dataset_id"]],
            "code": (
                'daily = datasets["dataset_1"].copy()\n'
                'result = daily.groupby("ts_code")["pct_chg"].sum().reset_index()'
            ),
        },
        lambda stage, message: progress.append((stage, message)),
    )

    assert payload["row_count"] == 1
    row = payload["preview"][0]
    assert row["ts_code"] == "000001.SZ"
    assert row["period_return_pct"] == pytest.approx(6.1106)
    assert sandbox.calls[0][0].startswith('daily = datasets["dataset_1"]')
    assert set(sandbox.calls[0][1]) == {"dataset_1"}
    assert "rank_recent_market_return" not in {
        definition["function"]["name"] for definition in toolbox.definitions
    }
    assert {stage for stage, _message in progress} == {"querying", "calculating"}


def test_model_preview_keeps_top_ten_rows_and_retains_the_complete_dataset():
    rows = [
        {"ts_code": f"{index:06d}.SZ", "close": float(index)}
        for index in range(12)
    ]
    result = QueryResult(
        query_id="ranking",
        provider="tushare",
        operation="daily",
        status=QueryStatus.SUCCESS,
        columns=["ts_code", "close"],
        rows=rows,
        row_count=len(rows),
    )

    toolbox = ResearchToolbox(FakeProvider(), "request-1", session_dataset=result)
    payload = toolbox.call(
        "inspect_session_dataset",
        {},
        lambda _stage, _message: None,
    )

    assert len(payload["preview"]) == 10
    assert payload["preview_truncated"] is True
    assert payload["dataset_scope"] == "complete_retained_result"


def test_toolbox_restores_session_dataset_and_archives_follow_up_result(tmp_path):
    session_result = QueryResult(
        query_id="previous_final",
        provider="tushare",
        operation="daily_basic",
        status=QueryStatus.SUCCESS,
        columns=["ts_code", "total_mv"],
        rows=[
            {"ts_code": "000001.SZ", "total_mv": 120.0},
            {"ts_code": "600000.SH", "total_mv": 90.0},
        ],
        row_count=2,
        completeness="complete",
    )
    store = MemoryAnalysisTaskStore()
    toolbox = ResearchToolbox(
        FakeProvider(),
        "request-1",
        artifact_dir=tmp_path,
        dataset_archive=store,
        task_id="agent-task",
        session_dataset=session_result,
    )

    inspected = toolbox.call(
        "inspect_session_dataset",
        {},
        lambda _stage, _message: None,
    )
    ranked = toolbox.call(
        "rank_dataset",
        {
            "dataset_id": "session_dataset",
            "sort_by": "total_mv",
            "direction": "asc",
            "limit": 1,
            "fields": ["ts_code", "total_mv"],
        },
        lambda _stage, _message: None,
    )
    toolbox.call(
        "export_excel",
        {
            "dataset_id": ranked["dataset_id"],
            "title": "Filtered result",
            "methodology": "Continued from the complete session dataset.",
            "column_notes": [
                {"column": "ts_code", "note": "A股证券代码，含交易所后缀。"},
                {"column": "total_mv", "note": "当日总市值，单位为万元。"},
            ],
        },
        lambda _stage, _message: None,
    )

    assert inspected["dataset_id"] == "session_dataset"
    assert inspected["row_count"] == 2
    assert ranked["preview"] == [{"ts_code": "600000.SH", "total_mv": 90.0}]
    assert store.promote_session_workspace("chat:session", "agent-task") is True
    restored = store.get_session_workspace("chat:session")
    assert restored is not None
    assert restored.rows == ranked["preview"]
    manifest = store.get_research_manifest("agent-task")
    assert manifest is not None
    assert manifest.final_dataset_id == ranked["dataset_id"]
    assert [evidence.dataset_id for evidence in manifest.datasets] == [
        "session_dataset",
        ranked["dataset_id"],
    ]


def test_toolbox_persists_machine_readable_dataset_evidence_without_raw_code():
    class UnitProvider(FakeProvider):
        def validate_query(self, operation, params, fields):
            assert operation == "daily"
            assert params == {"trade_date": "20260916"}
            assert fields == ["ts_code", "amount"]

        def query(self, operation, params, fields, **_kwargs):
            self.validate_query(operation, params, fields)
            return pd.DataFrame(
                [
                    {"ts_code": "000001.SZ", "amount": 860_000.0},
                    {"ts_code": "600000.SH", "amount": None},
                ]
            )

    class EchoSandbox:
        def run(self, _code, datasets):
            return datasets["dataset_1"].copy()

    store = MemoryAnalysisTaskStore()
    toolbox = ResearchToolbox(
        UnitProvider(),
        "request-1",
        python_sandbox=EchoSandbox(),
        dataset_archive=store,
        task_id="agent-task",
    )
    queried = toolbox.call(
        "query_market_data",
        {
            "operation": "daily",
            "params": {"trade_date": "20260916"},
            "fields": ["ts_code", "amount"],
        },
        lambda _stage, _message: None,
    )
    raw_code = 'result = datasets["dataset_1"].copy()  # private analysis logic'
    analyzed = toolbox.call(
        "run_python_analysis",
        {
            "dataset_ids": [queried["dataset_id"]],
            "code": raw_code,
        },
        lambda _stage, _message: None,
    )

    manifest = store.get_research_manifest("agent-task")

    assert manifest is not None
    assert len(manifest.datasets) == 2
    query_evidence, python_evidence = manifest.datasets
    assert query_evidence.operation_parameters == {"trade_date": "20260916"}
    assert query_evidence.requested_fields == ["ts_code", "amount"]
    assert query_evidence.row_count == 2
    assert query_evidence.missing_value_counts == {"ts_code": 0, "amount": 1}
    assert query_evidence.missing_value_rates == {"ts_code": 0.0, "amount": 0.5}
    assert query_evidence.field_units == {"amount": "thousands of CNY (千元)"}
    assert python_evidence.dataset_id == analyzed["dataset_id"]
    assert python_evidence.source_dataset_ids == [queried["dataset_id"]]
    assert len(python_evidence.operation_parameters["code_sha256"]) == 64
    assert raw_code not in manifest.model_dump_json()


def test_transform_tool_exposes_complete_pipeline_contract():
    toolbox = ResearchToolbox(RecentReturnProvider(), "request-1")

    transform = next(
        definition["function"]
        for definition in toolbox.definitions
        if definition["function"]["name"] == "transform_dataset"
    )
    parameters = transform["parameters"]
    pipeline = parameters["properties"]["pipeline"]

    assert parameters["additionalProperties"] is False
    assert parameters["required"] == ["pipeline"]
    assert pipeline["properties"]["steps"]["items"] == {
        "$ref": "#/$defs/ResultPipelineStep"
    }
    assert "ResultPipelineStep" in parameters["$defs"]
    assert "operation" in parameters["$defs"]["ResultPipelineStep"]["properties"]


def test_research_workbook_preserves_numeric_values_and_source_context():
    result = QueryResult(
        query_id="ranked",
        provider="tushare",
        operation="daily_basic",
        status=QueryStatus.SUCCESS,
        columns=["ts_code", "pe", "dividend_yield"],
        rows=[{"ts_code": "600519.SH", "pe": 24.5, "dividend_yield": 0.021}],
        row_count=1,
        completeness="complete",
    )

    path = build_research_workbook(
        result,
        "Valuation results",
        "Ranked by PE.",
        column_notes={
            "ts_code": "A股证券代码，含交易所后缀。",
            "pe": "市盈率 = 收盘价 / 每股收益。",
            "dividend_yield": "股息率 = 近12个月每股分红 / 收盘价。",
        },
    )

    workbook = load_workbook(path, data_only=False)
    results = workbook["Results"]
    methodology = workbook["Methodology"]
    notes = workbook["列说明"]
    assert results["B6"].value == 24.5
    assert isinstance(results["B6"].value, float)
    assert methodology["B4"].value == "tushare"
    assert methodology["B5"].value == "daily_basic"
    assert methodology["B7"].value == "Ranked by PE."
    assert results["C6"].number_format == "0.00%"
    assert notes["A5"].value == "ts_code"
    assert notes["B5"].value == "A股证券代码，含交易所后缀。"
    assert notes["A6"].value == "pe"
    assert notes["B7"].value == "股息率 = 近12个月每股分红 / 收盘价。"
    assert results["A5"].comment is not None
    assert "证券代码" in results["A5"].comment.text
    assert all(
        cell.data_type != "e"
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )


def test_research_workbook_links_security_codes_to_quote_pages():
    result = QueryResult(
        query_id="events",
        provider="tushare",
        operation="event_study",
        status=QueryStatus.SUCCESS,
        columns=["signal_date", "ts_code", "hk_code"],
        rows=[
            {"signal_date": 20260401, "ts_code": "688220.SH", "hk_code": "00700.HK"},
            {"signal_date": 20260402, "ts_code": "000002.SZ", "hk_code": "2.HK"},
        ],
        row_count=2,
        completeness="complete",
    )

    path = build_research_workbook(result, "Signals", "Event study basis.")

    results = load_workbook(path)["Results"]
    assert results["A6"].hyperlink is None
    assert results["B6"].hyperlink.target == "https://stockpage.10jqka.com.cn/688220/"
    assert results["B7"].hyperlink.target == "https://stockpage.10jqka.com.cn/000002/"
    assert results["C6"].hyperlink.target == "https://stockpage.10jqka.com.cn/HK0700/"
    assert results["C7"].hyperlink.target == "https://stockpage.10jqka.com.cn/HK0002/"


def test_export_tool_rejects_incomplete_column_notes(tmp_path):
    toolbox = ResearchToolbox(FakeProvider(), "request-1", artifact_dir=tmp_path)
    queried = toolbox.call(
        "query_market_data",
        {
            "operation": "daily",
            "params": {"trade_date": "20260916"},
            "fields": ["ts_code", "close"],
        },
        lambda _stage, _message: None,
    )

    missing = pytest.raises(
        ValueError,
        match=r"column_notes must exactly cover.*missing notes for: close",
    )
    with missing:
        toolbox.call(
            "export_excel",
            {
                "dataset_id": queried["dataset_id"],
                "title": "Incomplete",
                "methodology": "Basis.",
                "column_notes": [
                    {"column": "ts_code", "note": "A股证券代码，含交易所后缀。"}
                ],
            },
            lambda _stage, _message: None,
        )
    with pytest.raises(ValueError, match="notes for unknown columns: name"):
        toolbox.call(
            "export_excel",
            {
                "dataset_id": queried["dataset_id"],
                "title": "Unexpected",
                "methodology": "Basis.",
                "column_notes": [
                    {"column": "ts_code", "note": "A股证券代码，含交易所后缀。"},
                    {"column": "close", "note": "未复权收盘价，单位为元。"},
                    {"column": "name", "note": "不在导出数据集中的列。"},
                ],
            },
            lambda _stage, _message: None,
        )
    with pytest.raises(ValueError, match=r"8-400 characters"):
        toolbox.call(
            "export_excel",
            {
                "dataset_id": queried["dataset_id"],
                "title": "Too short",
                "methodology": "Basis.",
                "column_notes": [
                    {"column": "ts_code", "note": "太短"},
                    {"column": "close", "note": "未复权收盘价，单位为元。"},
                ],
            },
            lambda _stage, _message: None,
        )


def test_agent_coordinator_reports_progress_answer_and_file(tmp_path):
    store = MemoryAnalysisTaskStore()
    dispatcher = RecordingDispatcher()
    coordinator = FeishuAgentCoordinator(store, dispatcher)
    request = FeishuAgentRequest(
        prompt="Research banks and export Excel.",
        conversation_id="tenant:chat:root:user",
        source_message_id="message-1",
    )
    task = coordinator.submit(request, task_id="agent-task")
    artifact_path = tmp_path / "result.xlsx"
    artifact_path.write_bytes(b"xlsx")

    class FakeRuntime:
        def run(self, received_request, progress):
            assert received_request == request
            progress("querying", "正在查询市场数据…")
            progress("querying", "正在查询市场数据…")
            return FeishuAgentOutcome(
                answer="研究完成。",
                artifact_path=artifact_path,
            )

    sink = RecordingSink()
    completed = coordinator.run(task.task_id, FakeRuntime(), sink)

    assert dispatcher.task_ids == ["agent-task"]
    assert completed.status == AnalysisTaskStatus.SUCCEEDED
    assert completed.answer == "研究完成。"
    assert sink.messages == [
        ("message-1", "正在理解问题并选择研究工具…"),
    ]
    assert sink.updates == [
        ("progress-message-1", "正在查询市场数据…"),
        ("progress-message-1", "研究完成。"),
    ]
    assert sink.files == [("message-1", artifact_path)]
    assert store.get_artifact("agent-task", "result.xlsx") == b"xlsx"
    assert isinstance(store.get("agent-task"), FeishuAgentTask)
    assert completed.execution_attempt_count == 1
    assert completed.execution_lease_owner is None
    assert completed.research_manifest is not None
    assert completed.result_fingerprint == completed.research_manifest.output_fingerprint
    assert len(completed.result_fingerprint) == 64
    assert {record.delivery_id: record.status for record in completed.delivery_outbox} == {
        "terminal-answer": "sent",
        "terminal-artifact": "sent",
    }


def test_agent_coordinator_rejects_a_duplicate_worker_during_execution():
    store = MemoryAnalysisTaskStore()
    coordinator = FeishuAgentCoordinator(store, RecordingDispatcher())
    task = coordinator.submit(
        FeishuAgentRequest(
            prompt="Research without duplicate execution.",
            conversation_id="tenant:chat:root:user",
            source_message_id="message-1",
        ),
        task_id="agent-task",
    )
    duplicate_sink = RecordingSink()

    class NeverRuntime:
        def run(self, _request, _progress):
            raise AssertionError("A duplicate worker must not execute research.")

    class ReentrantRuntime:
        calls = 0

        def run(self, _request, _progress):
            self.calls += 1
            duplicate = coordinator.run(task.task_id, NeverRuntime(), duplicate_sink)
            assert duplicate.status == AnalysisTaskStatus.RUNNING
            return FeishuAgentOutcome(answer="One execution completed.")

    runtime = ReentrantRuntime()
    completed = coordinator.run(task.task_id, runtime, RecordingSink())

    assert completed.status == AnalysisTaskStatus.SUCCEEDED
    assert completed.execution_attempt_count == 1
    assert runtime.calls == 1
    assert duplicate_sink.messages == []
    assert duplicate_sink.updates == []
    assert duplicate_sink.files == []


def test_agent_coordinator_persists_success_with_complete_outbox(tmp_path):
    class ObservingStore(MemoryAnalysisTaskStore):
        def __init__(self):
            super().__init__()
            self.successful_writes = []

        def put_claimed_feishu_task(self, task, lease_owner, lease_expires_at):
            if task.status == AnalysisTaskStatus.SUCCEEDED:
                self.successful_writes.append(task.model_copy(deep=True))
            super().put_claimed_feishu_task(task, lease_owner, lease_expires_at)

    store = ObservingStore()
    coordinator = FeishuAgentCoordinator(store, RecordingDispatcher())
    task = coordinator.submit(
        FeishuAgentRequest(
            prompt="Persist every terminal intent with success.",
            conversation_id="tenant:chat:root:user",
            source_message_id="message-1",
        ),
        task_id="agent-task",
    )
    artifact_path = tmp_path / "result.xlsx"
    artifact_path.write_bytes(b"workbook")

    class Runtime:
        def run(self, _request, _progress):
            return FeishuAgentOutcome(answer="Research completed.", artifact_path=artifact_path)

    coordinator.run(task.task_id, Runtime(), RecordingSink())

    first_success = store.successful_writes[0]
    assert first_success.result_fingerprint is not None
    assert {record.delivery_id: record.status for record in first_success.delivery_outbox} == {
        "terminal-answer": "pending",
        "terminal-artifact": "pending",
    }


def test_agent_coordinator_retries_artifact_outbox_without_research(tmp_path):
    store = MemoryAnalysisTaskStore()
    coordinator = FeishuAgentCoordinator(store, RecordingDispatcher())
    task = coordinator.submit(
        FeishuAgentRequest(
            prompt="Research once and retry delivery only.",
            conversation_id="tenant:chat:root:user",
            source_message_id="message-1",
        ),
        task_id="agent-task",
    )
    artifact_path = tmp_path / "result.xlsx"
    artifact_path.write_bytes(b"durable-workbook")

    class CountingRuntime:
        calls = 0

        def run(self, _request, _progress):
            self.calls += 1
            return FeishuAgentOutcome(
                answer="Research completed.",
                artifact_path=artifact_path,
            )

    class FailOnceFileSink(RecordingSink):
        def __init__(self):
            super().__init__()
            self.file_attempts = 0
            self.file_payloads = []

        def reply_file(self, message_id, path):
            self.file_attempts += 1
            if self.file_attempts == 1:
                raise RuntimeError("temporary file delivery failure")
            self.file_payloads.append((message_id, path.read_bytes()))

    runtime = CountingRuntime()
    sink = FailOnceFileSink()
    first = coordinator.run(task.task_id, runtime, sink)
    first_fingerprint = first.result_fingerprint

    assert first.status == AnalysisTaskStatus.SUCCEEDED
    assert first.progress_message == "研究完成，但附件发送失败。"
    assert any(
        record.delivery_id == "terminal-artifact" and record.status == "failed"
        for record in first.delivery_outbox
    )

    recovered = coordinator.run(task.task_id, runtime, sink)
    unchanged = coordinator.run(task.task_id, runtime, sink)

    assert runtime.calls == 1
    assert sink.file_attempts == 2
    assert sink.file_payloads == [("message-1", b"durable-workbook")]
    assert recovered.progress_message == "研究完成。"
    assert recovered.result_fingerprint == first_fingerprint
    assert unchanged.execution_attempt_count == 2
    assert any(
        record.delivery_id == "terminal-artifact" and record.status == "sent"
        for record in recovered.delivery_outbox
    )


def test_agent_task_accepts_records_created_before_reliability_fields():
    parsed = FeishuAgentTask.model_validate(
        {
            "task_type": "feishu_agent",
            "task_id": "historical-task",
            "status": "queued",
            "request": {
                "prompt": "Historical request",
                "conversation_id": "conversation",
                "source_message_id": "message",
            },
            "created_at": datetime.now(timezone.utc),
            "updated_at": datetime.now(timezone.utc),
        }
    )

    assert parsed.execution_attempt_count == 0
    assert parsed.execution_lease_owner is None
    assert parsed.research_manifest is None
    assert parsed.result_fingerprint is None
    assert parsed.delivery_outbox == []


def test_agent_coordinator_preserves_success_when_file_delivery_fails(tmp_path):
    store = MemoryAnalysisTaskStore()
    coordinator = FeishuAgentCoordinator(store, RecordingDispatcher())
    request = FeishuAgentRequest(
        prompt="Research banks and export Excel.",
        conversation_id="tenant:chat:root:user",
        source_message_id="message-1",
    )
    task = coordinator.submit(request, task_id="agent-task")
    artifact_path = tmp_path / "result.xlsx"
    artifact_path.write_bytes(b"xlsx")

    class FakeRuntime:
        def run(self, _request, _progress):
            return FeishuAgentOutcome(
                answer="研究完成。",
                artifact_path=artifact_path,
            )

    class FailingFileSink(RecordingSink):
        def reply_file(self, _message_id, _path):
            raise RuntimeError("Feishu file upload failed.")

    sink = FailingFileSink()
    completed = coordinator.run(task.task_id, FakeRuntime(), sink)

    assert completed.status == AnalysisTaskStatus.SUCCEEDED
    assert completed.answer == "研究完成。"
    assert completed.progress_message == "研究完成，但附件发送失败。"
    assert completed.error is None
    assert sink.updates[-1] == (
        "progress-message-1",
        "研究完成。\n\n研究已完成，但附件发送失败。请稍后回复“重试”重新生成附件。",
    )


def test_agent_coordinator_publishes_tokenized_visualization_link(tmp_path):
    store = MemoryAnalysisTaskStore()
    coordinator = FeishuAgentCoordinator(
        store,
        RecordingDispatcher(),
        public_app_url="https://research.example/",
    )
    task = coordinator.submit(
        FeishuAgentRequest(
            prompt="Create an interactive valuation chart.",
            conversation_id="tenant:chat:root:user",
            source_message_id="message-1",
        ),
        task_id="agent-task",
    )
    artifact_path = tmp_path / "result.xlsx"
    artifact_path.write_bytes(b"workbook")
    visualization = FeishuResearchVisualization(
        title="Valuation ranking",
        columns=["name", "pe_ttm"],
        numeric_columns=["pe_ttm"],
        rows=[{"name": "Example", "pe_ttm": 12.5}],
        source_row_count=1,
        truncated=False,
        suggested_x="name",
        suggested_y="pe_ttm",
    )

    class FakeRuntime:
        def run(self, _request, _progress):
            return FeishuAgentOutcome(
                answer="研究完成。",
                artifact_path=artifact_path,
                visualization=visualization,
            )

    sink = RecordingSink()
    completed = coordinator.run(task.task_id, FakeRuntime(), sink)

    terminal_message = sink.updates[-1][1]
    assert "研究结果页面（30天内有效，点击后直接查看）" in terminal_message
    assert "交互图表" not in terminal_message
    link = terminal_message.rsplit("\n", 1)[-1]
    token = link.split("token=", 1)[1]
    assert link.startswith("https://research.example/research/agent-task?")
    assert completed.visualization == visualization
    assert completed.visualization_token_hash == research_visualization_token_hash(
        token
    )
    assert token not in completed.model_dump_json()
    assert completed.visualization_expires_at is not None
    assert store.get_artifact("agent-task", "result.xlsx") == b"workbook"


def test_agent_coordinator_reserves_final_message_update():
    store = MemoryAnalysisTaskStore()
    coordinator = FeishuAgentCoordinator(store, RecordingDispatcher())
    task = coordinator.submit(
        FeishuAgentRequest(
            prompt="Research a broad market question.",
            conversation_id="tenant:chat:root:user",
            source_message_id="message-1",
        ),
        task_id="agent-task",
    )

    class VerboseRuntime:
        def run(self, _request, progress):
            for index in range(30):
                progress("querying", f"Progress {index}")
            return FeishuAgentOutcome(answer="Final answer.")

    sink = RecordingSink()
    completed = coordinator.run(task.task_id, VerboseRuntime(), sink)

    assert completed.status == AnalysisTaskStatus.SUCCEEDED
    assert len(sink.messages) == 1
    assert len(sink.updates) == 20
    assert sink.updates[-1] == ("progress-message-1", "Final answer.")


def test_feishu_client_returns_visible_reply_message_id():
    session = SequenceSession(
        [
            FakeResponse({"code": 0, "tenant_access_token": "tenant-token"}),
            FakeResponse({"code": 0, "data": {"message_id": "reply-message-1"}}),
        ]
    )

    reply_message_id = FeishuOpenApiClient(
        "app-id", "app-secret", session=session
    ).reply(
        "message-1",
        "Research accepted.",
    )

    assert reply_message_id == "reply-message-1"
    assert session.calls[1][0].endswith("/im/v1/messages/message-1/reply")
    assert "params" not in session.calls[1][1]
    assert session.calls[1][1]["json"]["msg_type"] == "text"
    assert "reply_in_thread" not in session.calls[1][1]["json"]


def test_feishu_client_updates_existing_text_reply():
    session = SequenceSession(
        [
            FakeResponse({"code": 0, "tenant_access_token": "tenant-token"}),
            FakeResponse({"code": 0}),
        ]
    )

    FeishuOpenApiClient("app-id", "app-secret", session=session).update(
        "reply-message-1",
        "Research completed.",
    )

    assert session.calls[1][0].endswith("/im/v1/messages/reply-message-1")
    assert session.calls[1][1]["json"]["msg_type"] == "text"
    assert session.calls[1][1]["json"]["content"] == (
        '{"text": "Research completed."}'
    )


def test_feishu_client_replies_with_interactive_card():
    session = SequenceSession(
        [
            FakeResponse({"code": 0, "tenant_access_token": "tenant-token"}),
            FakeResponse({"code": 0, "data": {"message_id": "card-message-1"}}),
        ]
    )
    card = {"elements": [{"tag": "div"}]}

    message_id = FeishuOpenApiClient(
        "app-id", "app-secret", session=session
    ).reply_card("message-1", card)

    assert message_id == "card-message-1"
    assert session.calls[1][0].endswith("/im/v1/messages/message-1/reply")
    assert session.calls[1][1]["json"] == {
        "msg_type": "interactive",
        "content": '{"elements": [{"tag": "div"}]}',
    }


def test_feishu_client_uploads_and_replies_with_excel_file(tmp_path):
    path = tmp_path / "result.xlsx"
    path.write_bytes(b"workbook")
    session = SequenceSession(
        [
            FakeResponse({"code": 0, "tenant_access_token": "tenant-token"}),
            FakeResponse({"code": 0, "data": {"file_key": "file-key"}}),
            FakeResponse({"code": 0}),
        ]
    )

    FeishuOpenApiClient("app-id", "app-secret", session=session).reply_file(
        "message-1",
        path,
    )

    assert session.calls[1][0].endswith("/im/v1/files")
    assert session.calls[1][1]["data"] == {
        "file_type": "xls",
        "file_name": "result.xlsx",
    }
    assert session.calls[2][0].endswith("/im/v1/messages/message-1/reply")
    assert "params" not in session.calls[2][1]
    assert session.calls[2][1]["json"]["msg_type"] == "file"
    assert "reply_in_thread" not in session.calls[2][1]["json"]


def test_feishu_client_reports_file_upload_permission_error(tmp_path):
    path = tmp_path / "result.xlsx"
    path.write_bytes(b"workbook")
    session = SequenceSession(
        [
            FakeResponse({"code": 0, "tenant_access_token": "tenant-token"}),
            FakeResponse(
                {
                    "code": 99991672,
                    "msg": (
                        "Access denied. Required scope: "
                        "im:resource:upload."
                    ),
                },
                status_code=400,
            ),
        ]
    )

    with pytest.raises(RuntimeError, match="99991672.*im:resource:upload"):
        FeishuOpenApiClient(
            "app-id",
            "app-secret",
            session=session,
        ).reply_file("message-1", path)


@pytest.mark.live
@pytest.mark.skipif(
    os.getenv("RUN_LIVE_ANALYSIS") != "1",
    reason="Set RUN_LIVE_ANALYSIS=1 to call the configured model and Tushare.",
)
def test_live_feishu_agent_answers_one_historical_price_question():
    runtime = create_feishu_agent_runtime(Settings.from_env())

    outcome = runtime.run(
        FeishuAgentRequest(
            prompt=(
                "请查询贵州茅台（600519.SH）在2026年9月15日的收盘价，"
                "只回答代码、日期和收盘价。"
            ),
            conversation_id="live:feishu:agent:session",
            source_message_id="live-message",
        ),
        lambda _stage, _message: None,
    )

    assert "600519.SH" in outcome.answer
    assert "2026" in outcome.answer


@pytest.mark.live
@pytest.mark.skipif(
    os.getenv("RUN_LIVE_ANALYSIS") != "1",
    reason="Set RUN_LIVE_ANALYSIS=1 to call the configured model and Tushare.",
)
def test_live_feishu_agent_answers_reported_five_day_return_ranking():
    runtime = create_feishu_agent_runtime(Settings.from_env())

    outcome = runtime.run(
        FeishuAgentRequest(
            prompt="查出最近5个交易日累计涨幅最大的个股",
            conversation_id="live:feishu:agent:reported-five-day-return",
            source_message_id="live-reported-recent-return",
        ),
        lambda _stage, _message: None,
    )

    assert "二连板" not in outcome.answer
    assert "第三日" not in outcome.answer
    assert "复权收益" not in outcome.answer
    assert "并非复权" in outcome.answer or "复权" not in outcome.answer
    assert any(character.isdigit() for character in outcome.answer)


@pytest.mark.live
@pytest.mark.skipif(
    os.getenv("RUN_LIVE_ANALYSIS") != "1",
    reason="Set RUN_LIVE_ANALYSIS=1 to call the configured model and Tushare.",
)
def test_live_feishu_agent_answers_reported_thirty_day_return_ranking():
    runtime = create_feishu_agent_runtime(Settings.from_env())

    outcome = runtime.run(
        FeishuAgentRequest(
            prompt="查出最近30个交易日累计涨幅最大的个股",
            conversation_id="live:feishu:agent:reported-thirty-day-return",
            source_message_id="live-reported-thirty-day-return",
        ),
        lambda _stage, _message: None,
    )

    assert "无法可靠查出" not in outcome.answer
    assert "工具调用额度" not in outcome.answer
    assert "只返回前 20 行" not in outcome.answer
    assert re.search(r"\b\d{6}\.(?:SH|SZ|BJ)\b", outcome.answer)
    assert "%" in outcome.answer


@pytest.mark.live
@pytest.mark.skipif(
    os.getenv("RUN_LIVE_ANALYSIS") != "1",
    reason="Set RUN_LIVE_ANALYSIS=1 to call the configured model and Tushare.",
)
def test_live_feishu_agent_resolves_bare_menu_number_from_history():
    """A bare menu reply must resolve through the retained prior turn."""
    runtime = create_feishu_agent_runtime(Settings.from_env())

    outcome = runtime.run(
        FeishuAgentRequest(
            prompt="2",
            conversation_id="live:feishu:agent:reported-menu-selection",
            source_message_id="live-reported-menu-selection",
            conversation=[
                FeishuAgentConversationTurn(
                    prompt="研究对象为A股所有股票",
                    answer=(
                        "已记录研究范围为A股全部股票。请问接下来您希望对这批股票"
                        "进行哪类分析？回复序号选择：1 涨跌表现统计 2 财务指标统计"
                        " 3 估值水平统计。"
                    ),
                )
            ],
        ),
        lambda _stage, _message: None,
    )

    assert "没有历史上下文" not in outcome.answer
    assert "无法确定" not in outcome.answer
    assert "财务" in outcome.answer


@pytest.mark.live
@pytest.mark.skipif(
    os.getenv("RUN_LIVE_ANALYSIS") != "1",
    reason="Set RUN_LIVE_ANALYSIS=1 to call the configured model and Tushare.",
)
def test_live_feishu_agent_clarifies_reported_ambiguous_pe_ranking():
    runtime = create_feishu_agent_runtime(Settings.from_env())

    outcome = runtime.run(
        FeishuAgentRequest(
            prompt="你能查到今天市盈率前10的股票吗",
            conversation_id="live:feishu:agent:ambiguous-pe-ranking",
            source_message_id="live-ambiguous-pe-ranking",
        ),
        lambda _stage, _message: None,
    )

    assert "市盈率" in outcome.answer
    assert "1." in outcome.answer
    assert "推荐" in outcome.answer
    assert "请回复序号" in outcome.answer


@pytest.mark.live
@pytest.mark.skipif(
    os.getenv("RUN_LIVE_ANALYSIS") != "1",
    reason="Set RUN_LIVE_ANALYSIS=1 to call the configured model and Tushare.",
)
def test_live_feishu_agent_answers_reported_precise_pe_ranking():
    runtime = create_feishu_agent_runtime(Settings.from_env())

    outcome = runtime.run(
        FeishuAgentRequest(
            prompt=(
                "查询最近一个已完成交易日，全A股中 PE_TTM 最低且大于0的10只股票，"
                "排除ST、退市整理和市盈率为空的股票，列出代码、名称、PE_TTM和总市值。"
            ),
            conversation_id="live:feishu:agent:precise-pe-ranking",
            source_message_id="live-precise-pe-ranking",
        ),
        lambda _stage, _message: None,
    )

    assert "DSML" not in outcome.answer
    assert "<｜｜" not in outcome.answer
    answered_with_data = "PE" in outcome.answer and any(
        character.isdigit() for character in outcome.answer
    )
    requested_clarification = (
        "推荐" in outcome.answer and "请回复序号" in outcome.answer
    )
    assert answered_with_data or requested_clarification


def test_research_workbook_renders_compact_calendar_dates_not_grouped_numbers():
    from china_a_share.feishu_agent import compact_calendar_date_text

    assert compact_calendar_date_text("最近大涨日", 20260917) == "2026-09-17"
    assert compact_calendar_date_text("trade_date", "20260917 ") == "2026-09-17"
    assert compact_calendar_date_text("上市时间", 20260917) == "2026-09-17"
    assert compact_calendar_date_text("上市日期", 20261345) == 20261345
    assert compact_calendar_date_text("市值_元", 20260917) == 20260917
    assert compact_calendar_date_text("最近大涨日", 708) == 708

    result = QueryResult(
        query_id="ranked",
        provider="tushare",
        operation="daily",
        status=QueryStatus.SUCCESS,
        columns=["代码", "名称", "最近大涨日", "市值_元"],
        rows=[
            {
                "代码": "002487.SZ",
                "名称": "大金重工",
                "最近大涨日": 20260917,
                "市值_元": 20260917,
            }
        ],
        row_count=1,
        completeness="complete",
    )

    path = build_research_workbook(
        result,
        "大涨日触发统计",
        "未复权日线口径。",
        column_notes={
            "代码": "A股证券代码，含交易所后缀。",
            "名称": "证券简称。",
            "最近大涨日": "窗口内最近一次收盘涨幅达到5%的交易日。",
            "市值_元": "收盘总市值，单位为元。",
        },
    )

    results = load_workbook(path, data_only=False)["Results"]
    assert results["C6"].value == "2026-09-17"
    assert results["D6"].value == 20260917
